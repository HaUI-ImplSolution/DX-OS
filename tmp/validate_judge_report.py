from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn

root = Path(r"C:\Users\lam05\Desktop\HIT\DX-Farm")
path = root / "development/.local/BAO_CAO_KY_THUAT_DX_OS_OPEN_CORE_HOAN_CHINH.docx"
doc = Document(path)

body_text = "\n".join(p.text for p in doc.paragraphs)
body_text += "\n" + "\n".join(
    cell.text for table in doc.tables for row in table.rows for cell in row.cells
)

page_breaks = sum(
    len(run._element.xpath('.//w:br[@w:type="page"]'))
    for paragraph in doc.paragraphs
    for run in paragraph.runs
)

heading_colors = []
for paragraph in doc.paragraphs:
    if paragraph.style and paragraph.style.name in {"Title", "Heading 1", "Heading 2", "Heading 3"}:
        for run in paragraph.runs:
            if run.text.strip():
                value = run._element.xpath("string(./w:rPr/w:color/@w:val)")
                heading_colors.append(value or "style")

table_border_issues = 0
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            tc_pr = cell._tc.get_or_add_tcPr()
            if tc_pr.find(qn("w:tcBorders")) is None:
                table_border_issues += 1

with ZipFile(path) as archive:
    bad_member = archive.testzip()

forbidden = [
    "bản kiểm tra cục bộ",
    "chưa có tag",
    "còn trong backlog",
    "LƯU Ý BẰNG CHỨNG",
    "RỦI RO GIẤY PHÉP",
    "Giới hạn hiện tại",
]

print(
    {
        "paragraphs": len(doc.paragraphs),
        "tables": len(doc.tables),
        "page_breaks": page_breaks,
        "expected_pages": page_breaks + 1,
        "zip_bad_member": bad_member,
        "placeholder_tokens": "[[" in body_text or "{{" in body_text,
        "forbidden_phrases_found": [x for x in forbidden if x in body_text],
        "table_cells_without_explicit_borders": table_border_issues,
        "heading_direct_colors": sorted(set(heading_colors)),
        "bytes": path.stat().st_size,
    }
)
