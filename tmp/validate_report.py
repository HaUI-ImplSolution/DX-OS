from pathlib import Path
from zipfile import ZipFile

from docx import Document
from pypdf import PdfReader

root = Path(r"C:\Users\lam05\Desktop\HIT\DX-Farm")
docx_path = root / "development/docs/submission/BAO_CAO_KY_THUAT_DX_OS_OPEN_CORE.docx"
pdf_path = root / "development/docs/submission/BAO_CAO_KY_THUAT_DX_OS_OPEN_CORE.pdf"

doc = Document(docx_path)
text = "\n".join(p.text for p in doc.paragraphs)
text += "\n" + "\n".join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
page_breaks = sum(
    1
    for paragraph in doc.paragraphs
    for run in paragraph.runs
    for node in run._element.xpath('.//w:br[@w:type="page"]')
)
with ZipFile(docx_path) as archive:
    bad_member = archive.testzip()

reader = PdfReader(pdf_path)
pdf_text = "".join((page.extract_text() or "") for page in reader.pages)

print(
    {
        "docx_paragraphs": len(doc.paragraphs),
        "docx_tables": len(doc.tables),
        "docx_page_breaks": page_breaks,
        "docx_zip_bad_member": bad_member,
        "docx_has_placeholder_tokens": "[[" in text or "{{" in text,
        "pdf_pages": len(reader.pages),
        "pdf_replacement_chars": pdf_text.count("\ufffd"),
        "docx_bytes": docx_path.stat().st_size,
        "pdf_bytes": pdf_path.stat().st_size,
    }
)
