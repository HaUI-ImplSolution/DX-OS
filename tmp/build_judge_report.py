from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(r"C:\Users\lam05\Desktop\HIT\DX-Farm")
OUTPUT = ROOT / "development" / ".local" / "BAO_CAO_KY_THUAT_DX_OS_OPEN_CORE_HOAN_CHINH.docx"

BLACK = "000000"
NAVY = "1F4E78"
PALE_BLUE = "F2F7FB"
PALE_GRAY = "F6F6F6"
BORDER = "D9D9D9"
MID_GRAY = "5E6670"
WHITE = "FFFFFF"


def set_font(run, name="Arial", size=None, bold=None, color=BLACK, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=110, start=110, bottom=110, end=110):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_borders(cell, color=BORDER, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tag = OxmlElement("w:tblHeader")
    tag.set(qn("w:val"), "true")
    tr_pr.append(tag)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Trang ")
    set_font(run, size=8, color=MID_GRAY)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.65)
    section.bottom_margin = Cm(1.55)
    section.left_margin = Cm(1.85)
    section.right_margin = Cm(1.85)
    section.header_distance = Cm(0.75)
    section.footer_distance = Cm(0.7)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    normal.font.size = Pt(10.2)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(5)

    title = doc.styles["Title"]
    title.font.name = "Arial"
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    title.font.size = Pt(25)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(BLACK)
    title.paragraph_format.space_after = Pt(10)

    for name, size, before, after in (
        ("Heading 1", 17, 0, 8),
        ("Heading 2", 12.5, 8, 4),
        ("Heading 3", 10.5, 6, 3),
    ):
        style = doc.styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.text = "DX OS OPEN CORE     BÁO CÁO KỸ THUẬT"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        set_font(run, size=7.5, bold=True, color=BLACK)

    footer = section.footer.paragraphs[0]
    add_page_field(footer)


def paragraph(doc, text="", *, bold_lead=None, align=None, size=None, italic=False, keep=False):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.widow_control = True
    p.paragraph_format.keep_together = keep
    if bold_lead and text.startswith(bold_lead):
        lead = p.add_run(bold_lead)
        set_font(lead, size=size, bold=True)
        rest = p.add_run(text[len(bold_lead):])
        set_font(rest, size=size, italic=italic)
    else:
        run = p.add_run(text)
        set_font(run, size=size, italic=italic)
    return p


def bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Cm(0.55)
        p.paragraph_format.first_line_indent = Cm(-0.25)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.08
        p.paragraph_format.widow_control = True
        run = p.add_run(item)
        set_font(run, size=9.7)


def numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.left_indent = Cm(0.65)
        p.paragraph_format.first_line_indent = Cm(-0.3)
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(item)
        set_font(run, size=9.7)


def table(doc, headers, rows, widths, font_size=8.5, center_columns=()):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    header = t.rows[0]
    repeat_header(header)
    prevent_row_split(header)
    for idx, label in enumerate(headers):
        cell = header.cells[idx]
        cell.width = Cm(widths[idx])
        shade_cell(cell, NAVY)
        set_cell_margins(cell, 120, 110, 120, 110)
        set_cell_borders(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        run = p.add_run(label)
        set_font(run, size=font_size, bold=True, color=WHITE)
    for row_index, row in enumerate(rows):
        cells = t.add_row().cells
        prevent_row_split(t.rows[-1])
        for idx, value in enumerate(row):
            cell = cells[idx]
            cell.width = Cm(widths[idx])
            shade_cell(cell, WHITE if row_index % 2 == 0 else PALE_BLUE)
            set_cell_margins(cell)
            set_cell_borders(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx in center_columns else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.03
            run = p.add_run(str(value))
            set_font(run, size=font_size)
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(0)
    return t


def page_break(doc):
    doc.add_page_break()


def section_title(doc, title):
    doc.add_paragraph(title, style="Heading 1")


def build():
    doc = Document()
    configure_document(doc)
    props = doc.core_properties
    props.title = "Báo cáo kỹ thuật DX OS Open Core cho DX Farm và DX Lab"
    props.subject = "Hồ sơ kỹ thuật dự thi PMNM HaUI 2026"
    props.author = "Nhóm phát triển DX OS"
    props.keywords = "DX OS Open Core, H P D I, DX Farm, DX Lab, PMNM HaUI 2026"

    # Page 1
    p = paragraph(doc, "ĐẠI HỌC CÔNG NGHIỆP HÀ NỘI", align=WD_ALIGN_PARAGRAPH.CENTER, size=10)
    p.paragraph_format.space_before = Pt(16)
    p.runs[0].bold = True
    p = paragraph(doc, "CUỘC THI PHẦN MỀM NGUỒN MỞ PMNM HaUI 2026", align=WD_ALIGN_PARAGRAPH.CENTER, size=10)
    p.runs[0].bold = True
    p.paragraph_format.space_after = Pt(26)

    title = doc.add_paragraph("Báo cáo kỹ thuật DX OS Open Core cho DX Farm và DX Lab", style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = paragraph(
        doc,
        "Nền tảng lõi dùng chung cho chuyển đổi số nông nghiệp và vận hành CNTT đại học",
        align=WD_ALIGN_PARAGRAPH.CENTER,
        size=12,
        italic=True,
    )
    subtitle.paragraph_format.space_after = Pt(20)

    doc.add_paragraph("Tóm tắt", style="Heading 2")
    paragraph(
        doc,
        "Nhóm chúng em xây dựng DX OS theo kiến trúc Open Core ba tầng. Tầng lõi cung cấp định danh và SSO, cổng API, dữ liệu, workflow, lưu trữ, thông báo và cổng AI dưới dạng dịch vụ headless. Không gian H P D I tổ chức các năng lực số dùng chung. Hai ứng dụng DX Farm và DX Lab khai thác cùng nền tảng để giải quyết hai miền nghiệp vụ khác nhau.",
    )
    paragraph(
        doc,
        "DX Farm số hóa chuỗi báo bệnh bằng ảnh, chẩn đoán có kỹ sư duyệt, khai phun và kiểm soát thu hoạch theo thời gian cách ly. DX Lab số hóa yêu cầu hỗ trợ CNTT theo SLA, mượn thiết bị, đặt phòng máy, tra cứu tri thức và thống kê vận hành. Cơ chế chọn ứng dụng bằng DXOS_APP cùng ranh giới module được kiểm soát khi biên dịch giúp Core giữ nguyên khi thay đổi ứng dụng nghiệp vụ.",
    )
    paragraph(
        doc,
        "Kết luận chính. Sản phẩm thể hiện đúng định hướng Open Core trước và ứng dụng sau. Core có vai trò vận hành thực tế, H P D I được chứng minh bằng các luồng đầu cuối, còn hồ sơ nguồn mở cung cấp giấy phép, SBOM, tài liệu dựng từ source và hướng dẫn đóng góp.",
        bold_lead="Kết luận chính.",
    )
    table(
        doc,
        ["Nội dung", "Minh chứng tiêu biểu"],
        [
            ("Open Core headless", "dx-core, Keycloak, APISIX, PostgreSQL, MinIO và các facade dùng chung"),
            ("H P D I", "Vai và workspace, workflow, dữ liệu và BI, CV, RAG và luật phân loại"),
            ("Ứng dụng", "DX Farm và DX Lab chạy trên cùng Core"),
            ("Nguồn mở", "LICENSE, LICENSES, SBOM, CHANGELOG, CONTRIBUTING và quy trình dựng lại"),
        ],
        [4.5, 12.5],
        font_size=8.7,
    )
    p = paragraph(doc, "Hà Nội, tháng 10 năm 2026", align=WD_ALIGN_PARAGRAPH.CENTER, size=9, italic=True)
    p.paragraph_format.space_before = Pt(10)

    # Page 2
    page_break(doc)
    section_title(doc, "1 Bài toán mục tiêu và phạm vi")
    paragraph(
        doc,
        "Đề thi yêu cầu đội thi xây dựng phần lõi Open Core trước, sau đó phát triển ứng dụng nghiệp vụ khai thác các dịch vụ của lõi. Sản phẩm vì vậy cần chứng minh được vai trò nền tảng của Core, mối liên hệ H P D I và khả năng tích hợp các thành phần nguồn mở thành một hệ thống thống nhất.",
    )
    doc.add_paragraph("1.1 Mục tiêu thiết kế", style="Heading 2")
    bullets(
        doc,
        [
            "Tách dịch vụ lõi khỏi giao diện và khỏi mô hình dữ liệu đặc thù của từng ứng dụng.",
            "Cung cấp một điểm truy cập có kiểm soát, xác thực tập trung và phân quyền theo vai cùng phạm vi dữ liệu.",
            "Chuẩn hóa workflow để lưu trạng thái, người xử lý, thời gian và lịch sử quyết định.",
            "Quản lý dữ liệu có cấu trúc trong PostgreSQL, tệp trong object storage và dữ liệu BI qua các view được cấp quyền.",
            "Đưa Intelligence vào quy trình có căn cứ, có cơ chế giải thích và có người chịu trách nhiệm quyết định.",
        ],
    )
    doc.add_paragraph("1.2 Hai ứng dụng minh họa", style="Heading 2")
    table(
        doc,
        ["Ứng dụng", "Người dùng", "Luồng nghiệp vụ chính", "Vai trò trong kiểm chứng"],
        [
            ("DX Farm", "Nông dân, kỹ sư, quản trị", "Báo bệnh bằng ảnh, chẩn đoán, duyệt, khai phun và kiểm soát PHI", "Chứng minh Core phục vụ mobile, web, AI, dữ liệu và workflow an toàn"),
            ("DX Lab", "Sinh viên, giảng viên, kỹ thuật viên, quản trị", "Yêu cầu hỗ trợ theo SLA, thiết bị, phòng máy, tri thức và thống kê", "Chứng minh cùng Core có thể phục vụ một miền nghiệp vụ đại học"),
        ],
        [2.3, 3.4, 6.1, 5.2],
        font_size=8.1,
    )
    doc.add_paragraph("1.3 Tiêu chí thành công", style="Heading 2")
    paragraph(
        doc,
        "Sản phẩm được đánh giá trên hai trục. Proof of FOSS thể hiện qua repository, giấy phép, dependency, SBOM, tài liệu và khả năng dựng lại. Chất lượng sản phẩm thể hiện qua kiến trúc ba tầng, độ ổn định của luồng demo, giá trị sử dụng, khả năng mở rộng và chất lượng giải thích kỹ thuật. Báo cáo tập trung vào những minh chứng có thể kiểm tra trực tiếp trong hệ thống và repository.",
    )
    paragraph(
        doc,
        "Phạm vi trình bày. Hai ứng dụng được sử dụng để kiểm chứng khả năng tái sử dụng của nền tảng. Trọng tâm của sản phẩm là Core headless và cách lớp H P D I cùng ứng dụng nghiệp vụ khai thác các dịch vụ dùng chung.",
        bold_lead="Phạm vi trình bày.",
    )

    # Page 3
    page_break(doc)
    section_title(doc, "2 Kiến trúc Open Core ba tầng")
    paragraph(
        doc,
        "Kiến trúc được tổ chức theo chiều từ dịch vụ nền lên năng lực số và ứng dụng nghiệp vụ. Mọi phụ thuộc từ ứng dụng hướng về hợp đồng của Core. Giao diện web và mobile nằm ở lớp năng lực hoặc ứng dụng, không nằm trong Tầng 1.",
    )
    table(
        doc,
        ["Tầng", "Vai trò", "Thành phần", "Minh chứng"],
        [
            ("Tầng 3 Ứng dụng", "Giải quyết bài toán cụ thể", "DX Farm và DX Lab", "Hai module nghiệp vụ và hai giao diện trên cùng một Core"),
            ("Tầng 2 H P D I", "Tổ chức năng lực vận hành số", "Human, Process, Data, Intelligence", "Vai và workspace, workflow, BI, CV, RAG và luật"),
            ("Tầng 1 Open Core", "Cung cấp dịch vụ headless", "Identity, Gateway, Data, Workflow, Storage, Notification và AI Gateway", "API và SPI dùng chung, không có giao diện người dùng cuối"),
        ],
        [3.0, 4.1, 5.5, 4.4],
        font_size=8.2,
    )
    doc.add_paragraph("2.1 Ranh giới module", style="Heading 2")
    table(
        doc,
        ["Module", "Trách nhiệm", "Quan hệ phụ thuộc"],
        [
            ("dx-core-api", "Interface, DTO, lỗi và hợp đồng mở rộng", "Không phụ thuộc ứng dụng"),
            ("dx-core", "Cài đặt dịch vụ dùng chung", "Chỉ phụ thuộc dx-core-api"),
            ("app-farm", "Miền nghiệp vụ nông nghiệp", "Chỉ phụ thuộc dx-core-api"),
            ("app-lab", "Miền nghiệp vụ vận hành CNTT", "Chỉ phụ thuộc dx-core-api"),
            ("dx-host", "Ghép Core và ứng dụng thành gói chạy", "Nạp app theo DXOS_APP"),
        ],
        [3.2, 8.0, 5.8],
        font_size=8.4,
    )
    paragraph(
        doc,
        "Ứng dụng không import implementation hoặc entity bên trong dx-core. Core cũng không import ứng dụng. Các facade như WorkflowOperations, ObjectStorage, UserContext và KnowledgeGateway cung cấp bề mặt dùng chung. Ràng buộc module giúp compiler phát hiện vi phạm kiến trúc ngay trong quá trình build.",
    )
    doc.add_paragraph("2.2 Cơ chế cắm ứng dụng", style="Heading 2")
    paragraph(
        doc,
        "Mỗi ứng dụng khai báo vai, loại workflow, notification, view BI, namespace tri thức, xử lý tệp và chỉ báo sức khỏe qua SPI. Khi DXOS_APP thay đổi, host chỉ nạp entity, migration, controller và lịch chạy của ứng dụng được chọn. Gateway dùng cấu hình chung kết hợp với route riêng, còn web đọc manifest ứng dụng để nạp đúng menu và màn hình.",
    )

    # Page 4
    page_break(doc)
    section_title(doc, "3 Dịch vụ dùng chung của Core")
    paragraph(
        doc,
        "Tầng 1 tập hợp các dịch vụ nền mà cả Farm và Lab cùng sử dụng. Mỗi dịch vụ có hợp đồng rõ ràng, điểm kiểm soát bảo mật và cơ chế để ứng dụng khai báo phần đặc thù.",
    )
    table(
        doc,
        ["Dịch vụ", "Cách cung cấp", "Cách ứng dụng sử dụng"],
        [
            ("Identity và SSO", "Keycloak realm dxos, OIDC, đồng bộ vai do app khai báo", "Farm và Lab dùng vai khác nhau trên cùng cơ chế xác thực và RBAC"),
            ("API Gateway", "APISIX kiểm JWT, định tuyến và giới hạn tốc độ", "Web, mobile và API đi qua một điểm truy cập; Backend kiểm JWT lần hai"),
            ("Dữ liệu", "PostgreSQL và Flyway cho dữ liệu cấu trúc, MinIO cho tệp", "Mỗi app có migration riêng; client nhận link tệp có thời hạn"),
            ("Workflow", "Case type cấu hình được, event chỉ thêm và hộp duyệt dùng chung", "Farm khai báo năm loại duyệt; Lab khai báo ticket, mượn thiết bị và đặt phòng"),
            ("Thông báo", "Publisher dùng chung và kênh chat do app khai báo", "Thông báo vận hành tách khỏi transaction chính"),
            ("AI và tri thức", "Gateway nội bộ, circuit breaker và RAG theo namespace", "Farm và Lab dùng corpus cùng năng lực Intelligence riêng"),
        ],
        [3.0, 7.1, 6.9],
        font_size=7.9,
    )
    doc.add_paragraph("3.1 Luồng gọi của DX Farm", style="Heading 2")
    paragraph(
        doc,
        "Người dùng đăng nhập qua Keycloak và gửi báo cáo bệnh kèm ảnh qua APISIX. Backend kiểm tra token cùng phạm vi nông hộ, gọi AI để kiểm ảnh, làm sạch metadata rồi lưu tệp vào MinIO và bản ghi vào PostgreSQL. Khi kỹ sư yêu cầu chẩn đoán, Backend cấp link nội bộ có thời hạn cho AI, nhận kết quả, ghi Diagnosis và mở workflow case. Kỹ sư duyệt trước khi quy tắc xử lý được áp dụng.",
    )
    doc.add_paragraph("3.2 Luồng gọi của DX Lab", style="Heading 2")
    paragraph(
        doc,
        "Sinh viên nhập mô tả sự cố và nhận gợi ý loại, ưu tiên, lý do cùng bài tri thức liên quan. Khi tạo ticket, hệ thống tính SLA, mở workflow case và phát thông báo. Khóa phiên bản bảo đảm chỉ một kỹ thuật viên nhận thành công khi có thao tác đồng thời. Timeline ghi từng bước xử lý, còn ghi chú nội bộ được giới hạn cho người có vai phù hợp.",
    )
    paragraph(
        doc,
        "Nguyên tắc dữ liệu. Client chỉ biết Gateway. AI không truy cập trực tiếp cơ sở dữ liệu nghiệp vụ. PostgreSQL là nguồn sự thật, còn RAG phục vụ tra cứu và diễn đạt căn cứ.",
        bold_lead="Nguyên tắc dữ liệu.",
    )

    # Page 5
    page_break(doc)
    section_title(doc, "4 Không gian năng lực H P D I")
    paragraph(
        doc,
        "H P D I được thể hiện bằng hành vi vận hành và điểm kiểm soát trong hệ thống. Cùng một tập dịch vụ lõi tạo ra năng lực khác nhau tùy ứng dụng, vai trò và dữ liệu.",
    )
    table(
        doc,
        ["Không gian", "DX Farm", "DX Lab", "Minh chứng kỹ thuật"],
        [
            ("H Human", "Nông dân theo phạm vi hộ; kỹ sư và quản trị theo vai", "Sinh viên theo phạm vi ticket; kỹ thuật viên theo hàng đợi", "Keycloak, workspace, UserContext, RBAC và policy theo client"),
            ("P Process", "Duyệt thửa, vụ, chẩn đoán và khai phun; kiểm soát PHI", "Ticket theo SLA; duyệt mượn thiết bị; đặt phòng chống trùng", "Workflow case và event, state transition, optimistic locking và constraint DB"),
            ("D Data", "Bản ghi được duyệt, thời tiết, ảnh và view BI", "SLA, tải kỹ thuật viên, mức dùng phòng và tệp", "PostgreSQL, MinIO, Superset, metadata và provenance"),
            ("I Intelligence", "Kiểm ảnh, nhận diện bệnh, RAG và diễn đạt", "Luật phân loại giải thích được và RAG quy chế", "Ngưỡng, lý do, nguồn trích dẫn, human in the loop và fallback"),
        ],
        [2.0, 4.7, 4.7, 5.6],
        font_size=7.8,
    )
    doc.add_paragraph("4.1 Process và kiểm soát sai sót", style="Heading 2")
    paragraph(
        doc,
        "Các quy tắc an toàn được thực thi tại Backend hoặc cơ sở dữ liệu. Trong Farm, hệ thống kiểm soát thu hoạch theo PHI và chỉ cho phép dùng treatment rule đã được phê duyệt. Trong Lab, ràng buộc EXCLUDE USING gist chặn lịch phòng trùng giờ ngay cả khi hai yêu cầu đến đồng thời. Các cơ chế này bảo vệ quy trình độc lập với giao diện người dùng.",
    )
    doc.add_paragraph("4.2 Data và khả năng kiểm chứng", style="Heading 2")
    paragraph(
        doc,
        "Dữ liệu Farm được sinh từ bộ dữ liệu cấu trúc có trạng thái phê duyệt. Tài liệu RAG tách namespace cho Farm và Lab, kèm nguồn cùng giấy phép. BI chỉ đọc các view được cấp quyền. Tích hợp thời tiết có snapshot dự phòng và cờ nhận biết nguồn dữ liệu, giúp người dùng phân biệt dữ liệu trực tuyến với dữ liệu dự phòng.",
    )
    paragraph(
        doc,
        "Giá trị khác biệt. Sản phẩm kết hợp phân quyền theo phạm vi, workflow có lịch sử, ràng buộc đồng thời, dữ liệu có provenance và Intelligence giải thích được. Dashboard đóng vai trò quan sát, còn Core chịu trách nhiệm kiểm soát nghiệp vụ.",
        bold_lead="Giá trị khác biệt.",
    )

    # Page 6
    page_break(doc)
    section_title(doc, "5 Ứng dụng nghiệp vụ và kịch bản đầu cuối")
    doc.add_paragraph("5.1 DX Farm", style="Heading 2")
    table(
        doc,
        ["Bước", "Hành động", "Dịch vụ Core", "Kết quả"],
        [
            ("1", "Nông dân gửi ảnh báo bệnh", "SSO, Gateway, Storage và Data", "Ảnh được kiểm, làm sạch metadata và lưu riêng tư"),
            ("2", "AI phân tích", "AI Gateway, Knowledge và Storage", "Trả lớp bệnh, độ tin cậy và căn cứ"),
            ("3", "Kỹ sư duyệt", "Workflow và Notification", "Quyết định có người chịu trách nhiệm và lịch sử"),
            ("4", "Khai phun thuốc", "Structured Data và Workflow", "Rule đã duyệt cung cấp liều và PHI"),
            ("5", "Thu hoạch", "Poka yoke và Data", "Hệ thống kiểm tra điều kiện trước khi ghi nhận"),
        ],
        [1.2, 4.5, 5.1, 6.2],
        font_size=8.0,
        center_columns=(0,),
    )
    doc.add_paragraph("5.2 DX Lab", style="Heading 2")
    table(
        doc,
        ["Bước", "Hành động", "Dịch vụ Core", "Kết quả"],
        [
            ("1", "Sinh viên mô tả sự cố", "Rule Intelligence và Knowledge", "Gợi ý loại, ưu tiên, lý do và bài liên quan"),
            ("2", "Tạo ticket kèm ảnh", "Gateway, Storage, Data và Workflow", "Sinh mã, hạn SLA, case và timeline"),
            ("3", "Kỹ thuật viên nhận và xử lý", "Workflow, RBAC và locking", "Phân công nhất quán, trao đổi công khai và nội bộ"),
            ("4", "Hoàn tất và xác nhận", "Workflow và Notification", "Ghi chú, nguyên nhân gốc và đánh giá"),
            ("5", "Đặt phòng hoặc mượn thiết bị", "Workflow và DB constraints", "Duyệt theo vai và chặn xung đột"),
        ],
        [1.2, 4.5, 5.1, 6.2],
        font_size=8.0,
        center_columns=(0,),
    )
    doc.add_paragraph("5.3 Chuyển ứng dụng trên cùng nền tảng", style="Heading 2")
    paragraph(
        doc,
        "Lệnh khởi động với farm hoặc lab thay đổi DXOS_APP. Keycloak, PostgreSQL, MinIO và APISIX tiếp tục cung cấp hạ tầng chung. Backend nạp module ứng dụng được chọn, Gateway nhận route tương ứng, Superset đăng ký dataset và job RAG nạp corpus đúng namespace. Route của ứng dụng còn lại không xuất hiện trong runtime.",
    )

    # Page 7
    page_break(doc)
    section_title(doc, "6 Intelligence và quản trị dữ liệu")
    paragraph(
        doc,
        "Thiết kế Intelligence tuân theo nguyên tắc dữ kiện xác định trước, con người chịu trách nhiệm và kết quả có căn cứ. AI hỗ trợ nhận diện, tra cứu và ưu tiên nhưng không tự tạo số liệu an toàn hoặc quyết định nghiệp vụ cuối cùng.",
    )
    doc.add_paragraph("6.1 Pipeline Intelligence của DX Farm", style="Heading 2")
    table(
        doc,
        ["Giai đoạn", "Công nghệ và dữ liệu", "Cơ chế kiểm soát"],
        [
            ("Kiểm ảnh", "Kiểm chất lượng và crop classifier YOLO26s cls dạng ONNX", "Kiểm lớp lúa, độ tin cậy và chênh lệch với lớp khác"),
            ("Nhận diện bệnh", "YOLOv8n RiceLeafBD 3.1.0 cho năm bệnh lá lúa", "Confidence được trả về cùng kết quả để kỹ sư xem xét"),
            ("Tra cứu", "pgvector và embedding paraphrase multilingual MiniLM L12 v2", "Corpus có nguồn, giấy phép và namespace riêng"),
            ("Diễn đạt", "Qwen2.5 1.5B qua Ollama và LangGraph", "Không sinh số thuốc, liều hoặc PHI"),
            ("Quyết định", "Kỹ sư duyệt chẩn đoán và treatment rule", "Nhãn chốt được lưu có cấu trúc để kiểm chứng"),
        ],
        [2.8, 8.0, 6.2],
        font_size=8.1,
    )
    doc.add_paragraph("6.2 Intelligence của DX Lab", style="Heading 2")
    paragraph(
        doc,
        "Lab dùng luật từ khóa để gợi ý loại và ưu tiên. Kết quả kèm từ khóa khớp cùng lý do, giúp người dùng kiểm tra cách hệ thống suy luận. RAG bổ sung bài quy chế hoặc hướng dẫn liên quan. Người dùng vẫn quyết định nội dung ticket, còn kỹ thuật viên chịu trách nhiệm xử lý và kết luận.",
    )
    doc.add_paragraph("6.3 Quản trị dữ liệu và provenance", style="Heading 2")
    bullets(
        doc,
        [
            "PostgreSQL giữ dữ liệu nghiệp vụ và các quy tắc có cấu trúc.",
            "MinIO lưu ảnh cùng tệp đính kèm trong bucket riêng tư và cấp link có thời hạn.",
            "Tài liệu RAG ghi nguồn, giấy phép và namespace ứng dụng.",
            "Dữ liệu BI đi qua view chỉ đọc, tách khỏi bảng nghiệp vụ trực tiếp.",
            "Các nguồn bên ngoài có metadata và cơ chế fallback để duy trì luồng demo.",
        ],
    )

    # Page 8
    page_break(doc)
    section_title(doc, "7 Bảo mật kiểm thử và khả năng phục hồi")
    doc.add_paragraph("7.1 Bảo mật theo nhiều lớp", style="Heading 2")
    table(
        doc,
        ["Lớp kiểm soát", "Biện pháp"],
        [
            ("Biên hệ thống", "APISIX kiểm JWT, định tuyến theo app và giới hạn tốc độ; client chỉ truy cập qua Gateway"),
            ("Ứng dụng", "Backend xác thực JWT lần hai, áp dụng RBAC, phạm vi workspace và policy theo loại client"),
            ("Dữ liệu", "MinIO riêng tư, link có hạn, BI chỉ đọc view và schema được quản lý bằng Flyway"),
            ("Bí mật", "Cấu hình qua biến môi trường, env example và kiểm tra bắt buộc khi khởi động"),
            ("Tính toàn vẹn", "Optimistic locking, workflow event chỉ thêm và constraint chống xung đột"),
        ],
        [4.1, 12.9],
        font_size=8.4,
    )
    doc.add_paragraph("7.2 Kiểm thử và kiểm chứng", style="Heading 2")
    table(
        doc,
        ["Phạm vi", "Kết quả ghi nhận", "Ý nghĩa"],
        [
            ("Backend", "Bộ verify kiểm tra Core, ứng dụng và host", "Kiểm hợp đồng, workflow, quyền và tích hợp dữ liệu"),
            ("AI", "64 test cùng lint và dependency check trong đợt kiểm chứng CV", "Kiểm loader, policy, precheck, predict, diagnose và xử lý lỗi"),
            ("E2E Farm", "10 trên 10 bước", "Kiểm vai, duyệt, idempotency, ảnh, chẩn đoán và PHI"),
            ("E2E Lab", "16 trên 16 bước", "Kiểm ticket, SLA, thiết bị, phòng và tri thức"),
            ("Ranh giới app", "Route app không hoạt động không được nạp; app id được kiểm tra khi khởi động", "Kiểm lựa chọn ứng dụng và tính cô lập runtime"),
        ],
        [3.0, 7.0, 7.0],
        font_size=8.0,
    )
    doc.add_paragraph("7.3 Khả năng phục hồi", style="Heading 2")
    paragraph(
        doc,
        "Khi nguồn thời tiết tạm dừng, hệ thống dùng snapshot có đánh dấu. Khi LLM chưa sẵn sàng, RAG dùng nhánh evidence fallback. Gợi ý của Lab vẫn hoạt động bằng luật, thông báo chat chạy bất đồng bộ và mobile giữ hàng đợi gửi lại. Portal cùng endpoint trạng thái giúp người vận hành nhận biết tình trạng từng dịch vụ.",
    )
    paragraph(
        doc,
        "Phạm vi kiểm chứng. Kết quả trong báo cáo phản ánh các bài kiểm tra tự động và kịch bản đầu cuối đã ghi nhận trong tài liệu kiểm thử của dự án. Các số liệu phát hành được cập nhật theo workflow của phiên bản nộp bài.",
        bold_lead="Phạm vi kiểm chứng.",
    )

    # Page 9
    page_break(doc)
    section_title(doc, "8 Nguồn mở và khả năng tái lập")
    paragraph(
        doc,
        "Proof of FOSS được triển khai như một phần của kiến trúc sản phẩm. Repository lưu mã nguồn, lịch sử phát triển, giấy phép, dependency, SBOM, tài liệu vận hành và quy trình đóng góp. Môi trường có thể dựng lại bằng Docker Compose, Maven wrapper, lockfile và dữ liệu demo.",
    )
    doc.add_paragraph("8.1 Thành phần và giấy phép chính", style="Heading 2")
    table(
        doc,
        ["Nhóm", "Thành phần", "Giấy phép hoặc cách phân phối"],
        [
            ("Mã tự viết", "Core, backend ứng dụng, web, mobile và cấu hình hạ tầng", "Apache 2.0 và khai báo SPDX theo quy ước dự án"),
            ("AI", "Ultralytics và checkpoint CV", "Thư mục AI phân phối theo AGPL 3.0"),
            ("Định danh và Gateway", "Keycloak và Apache APISIX", "Apache 2.0"),
            ("Dữ liệu", "PostgreSQL, pgvector và MinIO", "PostgreSQL License và AGPL 3.0"),
            ("BI và Chat", "Apache Superset và Mattermost Team Edition", "Giấy phép được ghi nhận trong LICENSES"),
            ("Mô hình và dữ liệu", "Qwen, MiniLM, RiceLeafDiseaseBD và tài liệu RAG", "Ghi phiên bản, nguồn, vai trò và giấy phép trong metadata"),
        ],
        [3.1, 6.0, 7.9],
        font_size=7.9,
    )
    doc.add_paragraph("8.2 Dựng và chạy từ source", style="Heading 2")
    numbered(
        doc,
        [
            "Tạo cấu hình từ deploy .env.example và điền các biến bắt buộc.",
            "Chạy script start với farm hoặc lab để dựng stack, kiểm tra health và nạp dữ liệu demo.",
            "Dùng reset để tái lập môi trường sạch khi cần kiểm tra từ đầu.",
            "Truy cập Swagger, portal và các tài khoản demo theo hướng dẫn trong README.",
            "Chạy test module và E2E cho ứng dụng đang chọn.",
        ],
    )
    doc.add_paragraph("8.3 Phát triển cộng đồng", style="Heading 2")
    paragraph(
        doc,
        "CONTRIBUTING mô tả quy trình nhánh, pull request, kiểm tra theo module và cách thêm ứng dụng mới. Người đóng góp tạo module chỉ phụ thuộc dx-core-api, cài AppModule cùng các SPI cần thiết, thêm migration, route Gateway, dataset BI, corpus tri thức và giao diện. app-lab cung cấp một ví dụ hoàn chỉnh để đối chiếu.",
    )
    paragraph(
        doc,
        "Quản trị giấy phép. Các thành phần có nghĩa vụ phân phối khác nhau được công bố trong LICENSES. AI Service chạy như tiến trình riêng, giao tiếp với Backend qua HTTP và có giấy phép riêng cho thư mục AI.",
        bold_lead="Quản trị giấy phép.",
    )

    # Page 10
    page_break(doc)
    section_title(doc, "9 Kết quả và hướng phát triển")
    doc.add_paragraph("9.1 Đối chiếu yêu cầu đề thi", style="Heading 2")
    table(
        doc,
        ["Yêu cầu", "Kết quả", "Minh chứng chính"],
        [
            ("Core headless và ba tầng", "Đáp ứng", "Ranh giới module, SPI, Gateway, Identity, Data, Workflow và hai ứng dụng"),
            ("H P D I", "Đáp ứng", "Vai và phạm vi, workflow, PostgreSQL, MinIO, BI, CV, RAG và luật"),
            ("Ứng dụng thực tế", "Đáp ứng", "DX Farm và DX Lab có dữ liệu mẫu cùng kịch bản đầu cuối"),
            ("API và triển khai", "Đáp ứng", "REST OpenAPI, Docker Compose, script vận hành và dữ liệu demo"),
            ("Kiểm thử", "Đáp ứng phạm vi MVP", "Test module, E2E hai ứng dụng và kiểm tra ranh giới runtime"),
            ("Nguồn mở", "Đáp ứng hồ sơ kỹ thuật", "LICENSE, LICENSES, SBOM, CHANGELOG, CONTRIBUTING và tài liệu dựng lại"),
        ],
        [4.1, 3.1, 9.8],
        font_size=8.0,
        center_columns=(1,),
    )
    doc.add_paragraph("9.2 Hướng phát triển", style="Heading 2")
    paragraph(
        doc,
        "Kiến trúc hiện tại tạo nền tảng để mở rộng theo ba hướng. Thứ nhất, bổ sung ứng dụng nghiệp vụ mới qua dx-core-api và SPI mà không thay đổi Core. Thứ hai, nâng cao năng lực dữ liệu và Intelligence bằng quy trình đánh giá mô hình, provenance và quản trị tri thức. Thứ ba, mở rộng vận hành bằng cơ chế xử lý bất đồng bộ, quan sát hệ thống và triển khai nhiều bản dịch vụ khi quy mô sử dụng tăng.",
    )
    doc.add_paragraph("9.3 Kết luận", style="Heading 2")
    paragraph(
        doc,
        "DX OS thể hiện đúng tư duy nền tảng của đề thi. Core cung cấp dịch vụ dùng chung và giữ ranh giới ổn định khi chuyển từ DX Farm sang DX Lab. Không gian H P D I kết nối con người, quy trình, dữ liệu và trí tuệ trong các luồng nghiệp vụ có thể kiểm chứng. Kiến trúc module, cơ chế vận hành và hồ sơ nguồn mở cho phép sản phẩm tiếp tục phát triển như một nền tảng, thay vì dừng ở một ứng dụng đơn lẻ.",
    )
    doc.add_paragraph("Tài liệu tham chiếu", style="Heading 2")
    bullets(
        doc,
        [
            "Đề thi PMNM HaUI 2026 bản chính thức, các mục kiến trúc, nhiệm vụ và hồ sơ nộp bài.",
            "README, ARCHITECTURE, SUBMISSION và COUNCIL_QA của dự án.",
            "LICENSES, CONTRIBUTING và SBOM của các module.",
            "TESTING, README của backend, app-farm, app-lab và ai.",
        ],
    )
    p = paragraph(doc, "Nhóm phát triển DX OS Open Core", align=WD_ALIGN_PARAGRAPH.CENTER, size=10, italic=True)
    p.paragraph_format.space_before = Pt(8)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
