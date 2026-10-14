from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    BaseDocTemplate,
    Frame,
    KeepTogether,
    PageBreak,
    PageTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
)


ROOT = Path(r"C:\Users\lam05\Desktop\HIT\DX-Farm")
OUT_DIR = ROOT / "development" / "docs" / "submission"
DOCX_OUT = OUT_DIR / "BAO_CAO_KY_THUAT_DX_OS_OPEN_CORE.docx"
PDF_OUT = OUT_DIR / "BAO_CAO_KY_THUAT_DX_OS_OPEN_CORE.pdf"

NAVY = "17324D"
BLUE = "176B87"
TEAL = "1A8A7A"
PALE = "EAF3F6"
PALE_GREEN = "EAF6F2"
PALE_GOLD = "FFF4D6"
GRAY = "5D6A73"
LIGHT_GRAY = "F3F5F7"
WHITE = "FFFFFF"
RED = "A33A3A"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_text(cell, text, bold=False, color=NAVY, size=9.0, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(color)
    r.font.name = "Arial"
    r._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Arial")


def set_fixed_row_height(row, twips):
    tr_pr = row._tr.get_or_add_trPr()
    tr_height = OxmlElement("w:trHeight")
    tr_height.set(qn("w:val"), str(twips))
    tr_height.set(qn("w:hRule"), "atLeast")
    tr_pr.append(tr_height)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Trang ")
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def configure_doc(doc):
    sec = doc.sections[0]
    sec.page_width = Cm(21)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(1.55)
    sec.bottom_margin = Cm(1.5)
    sec.left_margin = Cm(1.75)
    sec.right_margin = Cm(1.75)
    sec.header_distance = Cm(0.7)
    sec.footer_distance = Cm(0.65)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    normal.font.size = Pt(10.1)
    normal.font.color.rgb = RGBColor.from_string("24313A")
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing = 1.12
    normal.paragraph_format.space_after = Pt(5)

    title = styles["Title"]
    title.font.name = "Arial"
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    title.font.size = Pt(27)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(NAVY)
    title.paragraph_format.space_after = Pt(10)

    for name, size, color, before, after in (
        ("Heading 1", 18, NAVY, 0, 8),
        ("Heading 2", 12.5, BLUE, 8, 4),
        ("Heading 3", 10.5, TEAL, 6, 3),
    ):
        st = styles[name]
        st.font.name = "Arial"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(color)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True

    for section in doc.sections:
        header = section.header
        p = header.paragraphs[0]
        p.text = "DX OS OPEN CORE   |   BÁO CÁO KỸ THUẬT"
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for r in p.runs:
            r.font.name = "Arial"
            r.font.size = Pt(7.5)
            r.font.bold = True
            r.font.color.rgb = RGBColor.from_string(BLUE)
        footer = section.footer
        fp = footer.paragraphs[0]
        add_page_number(fp)
        for r in fp.runs:
            r.font.name = "Arial"
            r.font.size = Pt(8)
            r.font.color.rgb = RGBColor.from_string(GRAY)


def add_doc_heading(doc, title, kicker=None):
    if kicker:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(kicker.upper())
        r.bold = True
        r.font.name = "Arial"
        r.font.size = Pt(8.5)
        r.font.color.rgb = RGBColor.from_string(TEAL)
    doc.add_paragraph(title, style="Heading 1")


def add_doc_para(doc, text, bold_lead=None, small=False):
    p = doc.add_paragraph()
    p.paragraph_format.widow_control = True
    if small:
        p.paragraph_format.space_after = Pt(3)
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead)
        r.bold = True
        r.font.color.rgb = RGBColor.from_string(NAVY)
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    if small:
        for r in p.runs:
            r.font.size = Pt(9.2)
    return p


def add_doc_bullets(doc, items, compact=True):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Cm(0.55)
        p.paragraph_format.first_line_indent = Cm(-0.25)
        p.paragraph_format.space_after = Pt(2.5 if compact else 4)
        p.paragraph_format.line_spacing = 1.05
        r = p.add_run(item)
        r.font.name = "Arial"
        r.font.size = Pt(9.5)


def add_doc_callout(doc, label, text, fill=PALE, accent=BLUE):
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Cm(3.0)
    table.columns[1].width = Cm(14.0)
    set_cell_shading(table.cell(0, 0), accent)
    set_cell_shading(table.cell(0, 1), fill)
    set_cell_text(table.cell(0, 0), label, bold=True, color=WHITE, size=9.2, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_text(table.cell(0, 1), text, color=NAVY, size=9.2)
    for c in table.rows[0].cells:
        set_cell_margins(c, top=120, bottom=120)
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_doc_table(doc, headers, rows, widths=None, font_size=8.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, h in enumerate(headers):
        set_cell_shading(hdr.cells[i], NAVY)
        set_cell_text(hdr.cells[i], h, bold=True, color=WHITE, size=font_size, align=WD_ALIGN_PARAGRAPH.CENTER)
        hdr.cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(hdr.cells[i])
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_shading(cells[i], WHITE if ridx % 2 == 0 else LIGHT_GRAY)
            set_cell_text(cells[i], str(value), color="24313A", size=font_size)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cells[i])
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Cm(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_doc_page_break(doc):
    doc.add_page_break()


def build_docx():
    doc = Document()
    configure_doc(doc)
    props = doc.core_properties
    props.title = "Báo cáo kỹ thuật DX OS Open Core cho DX Farm và DX Lab"
    props.subject = "Báo cáo kỹ thuật theo đề thi PMNM HaUI 2026"
    props.author = "Nhóm phát triển DX OS"
    props.keywords = "DX-OS, Open-Core, H-P-D-I, DX-Farm, DX-Lab, PMNM HaUI 2026"

    # Page 1
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(22)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("PMNM HaUI 2026")
    r.bold = True
    r.font.name = "Arial"
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor.from_string(TEAL)
    title = doc.add_paragraph("Báo cáo kỹ thuật DX OS Open Core cho DX Farm và DX Lab", style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.paragraph_format.space_after = Pt(18)
    rr = sub.add_run("Nền tảng lõi dùng chung cho chuyển đổi số nông nghiệp và vận hành CNTT đại học")
    rr.font.size = Pt(12)
    rr.font.color.rgb = RGBColor.from_string(GRAY)
    add_doc_callout(
        doc,
        "KẾT LUẬN CHÍNH",
        "Sản phẩm đáp ứng định hướng Open Core trước và ứng dụng sau: Tầng 1 cung cấp dịch vụ headless dùng chung; Tầng 2 tổ chức năng lực H P D I; Tầng 3 chứng minh khả năng tái sử dụng bằng hai ứng dụng DX Farm và DX Lab trên cùng một lõi.",
        fill=PALE_GREEN,
        accent=TEAL,
    )
    doc.add_paragraph("Tóm tắt điều hành", style="Heading 2")
    add_doc_para(doc, "DX OS được tổ chức như một nền tảng tích hợp thay vì một website nghiệp vụ đơn lẻ. Lõi dùng chung gồm định danh và SSO, cổng API, dữ liệu cấu trúc và phi cấu trúc, workflow, thông báo, lưu trữ, cổng AI và tri thức. Ứng dụng chỉ phụ thuộc hợp đồng trong dx-core-api; việc chọn ứng dụng thực thi qua biến DXOS_APP. Ranh giới này biến tính độc lập của Core thành một thuộc tính có thể kiểm tra khi biên dịch và khi chạy.")
    add_doc_para(doc, "DX Farm minh họa chuỗi báo bệnh bằng ảnh, chẩn đoán có người duyệt, khai phun và khóa thu hoạch theo thời gian cách ly. DX Lab minh họa hỗ trợ CNTT theo SLA, mượn thiết bị, đặt phòng máy, tri thức và thống kê. Hai miền nghiệp vụ khác nhau nhưng dùng cùng dịch vụ nền, qua đó chứng minh Core có vai trò thật và không bị khóa vào một giao diện cụ thể.")
    add_doc_table(doc, ["Trọng tâm đề thi", "Minh chứng trong sản phẩm"], [
        ("Open Core headless", "dx-core, Keycloak, APISIX, PostgreSQL, MinIO và các facade dùng chung"),
        ("Kiến trúc 3 tầng", "Core dùng chung, không gian H P D I, hai ứng dụng nghiệp vụ"),
        ("PoF nguồn mở", "License, bảng dependency, SBOM và quy trình dựng từ source; cần tạo tag và release nộp bài"),
        ("Intelligence có căn cứ", "CV và RAG ở Farm; luật giải thích được và RAG ở Lab; con người chốt quyết định"),
    ], widths=[4.7, 12.3], font_size=8.8)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    r = p.add_run("Ngày hoàn thiện báo cáo 13 tháng 10 năm 2026")
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor.from_string(GRAY)

    # Page 2
    add_doc_page_break(doc)
    add_doc_heading(doc, "1 Bài toán phạm vi và tiêu chí thành công", "Trang 2")
    add_doc_para(doc, "Đề thi yêu cầu xây dựng DX Lab HaUI theo tư duy Open Core trước, ứng dụng sau. Vì vậy, phạm vi sản phẩm không dừng ở một dashboard hoặc ứng dụng CRUD. Nền tảng phải cung cấp các dịch vụ dùng chung cho lớp H P D I và cho ít nhất một ứng dụng nghiệp vụ, đồng thời chứng minh được khả năng cài đặt, vận hành và phát triển bằng công nghệ nguồn mở.")
    doc.add_paragraph("1.1 Mục tiêu thiết kế", style="Heading 2")
    add_doc_bullets(doc, [
        "Tách phần lõi headless khỏi giao diện và khỏi mô hình dữ liệu đặc thù của từng ứng dụng.",
        "Cung cấp một cửa truy cập qua API Gateway, xác thực tập trung, phân quyền theo vai và phạm vi dữ liệu.",
        "Mô hình hóa quy trình dưới dạng case và sự kiện để lưu được trạng thái, người xử lý, thời gian và lịch sử thay đổi.",
        "Dùng PostgreSQL làm nguồn sự thật; object storage quản lý tệp; BI chỉ đọc các view đã kiểm soát.",
        "Đặt AI vào quy trình có căn cứ, có ngưỡng, có đường dự phòng và có con người chịu trách nhiệm quyết định.",
    ])
    doc.add_paragraph("1.2 Hai ứng dụng minh họa", style="Heading 2")
    add_doc_table(doc, ["Ứng dụng", "Người dùng chính", "Luồng xương sống", "Giá trị chứng minh"], [
        ("DX Farm", "Nông dân, kỹ sư, quản trị", "Ảnh bệnh lúa → AI → kỹ sư duyệt → khai phun → khóa PHI", "Core phục vụ mobile, web, AI, BI và kiểm soát an toàn nông nghiệp"),
        ("DX Lab", "Sinh viên, giảng viên, kỹ thuật viên, quản trị", "Gợi ý → tạo ticket → xử lý SLA → xác nhận; mượn thiết bị; đặt phòng", "Cùng Core phục vụ một miền nghiệp vụ đại học hoàn toàn khác"),
    ], widths=[2.3, 3.6, 6.1, 5.0], font_size=8.2)
    doc.add_paragraph("1.3 Tiêu chí thành công", style="Heading 2")
    add_doc_para(doc, "Thành công được đánh giá trên hai trục. Trục PoF yêu cầu repository công khai, lịch sử phát triển, giấy phép, release, dependency có nguồn gốc và quy trình dựng lại. Trục sản phẩm yêu cầu đúng kiến trúc, chạy ổn định, có giá trị sử dụng, có khả năng mở rộng và trình bày thuyết phục. Báo cáo này ưu tiên các minh chứng có thể kiểm tra thay vì mô tả giao diện.")
    add_doc_callout(doc, "PHẠM VI", "Hai ứng dụng được trình bày để chứng minh khả năng tái sử dụng. Trọng tâm chấm vẫn là nền tảng Open Core và cách ứng dụng khai thác các dịch vụ của lõi.", fill=PALE_GOLD, accent="C28A18")

    # Page 3
    add_doc_page_break(doc)
    add_doc_heading(doc, "2 Kiến trúc Open Core ba tầng", "Trang 3")
    add_doc_para(doc, "Kiến trúc đi từ dịch vụ lõi ở dưới lên năng lực số và ứng dụng nghiệp vụ. Luồng phụ thuộc luôn hướng về hợp đồng của Core; giao diện không nằm trong Tầng 1. Cách tổ chức này đáp ứng trực tiếp yêu cầu phân biệt nền tảng dùng chung với sản phẩm giao diện.")
    add_doc_table(doc, ["Tầng", "Vai trò", "Thành phần tiêu biểu", "Minh chứng"], [
        ("Tầng 3 Ứng dụng", "Giải quyết bài toán cụ thể", "DX Farm và DX Lab", "Hai module nghiệp vụ, hai giao diện, cùng một Core"),
        ("Tầng 2 H P D I", "Biến dịch vụ nền thành năng lực vận hành", "Human, Process, Data, Intelligence", "Vai và workspace; workflow; BI; CV, RAG và luật"),
        ("Tầng 1 Open Core", "Dịch vụ headless và PaaS cho ứng dụng", "Identity, Gateway, Data, Workflow, Storage, Notification, AI Gateway", "API và SPI dùng chung; không có UI người dùng cuối"),
    ], widths=[3.0, 4.2, 5.3, 4.5], font_size=8.4)
    doc.add_paragraph("2.1 Ranh giới module", style="Heading 2")
    add_doc_table(doc, ["Module", "Trách nhiệm", "Quan hệ phụ thuộc"], [
        ("dx-core-api", "Interface, DTO, lỗi và hợp đồng mở rộng", "Không phụ thuộc ứng dụng"),
        ("dx-core", "Cài đặt dịch vụ dùng chung", "Chỉ phụ thuộc dx-core-api"),
        ("app-farm", "Miền nghiệp vụ nông nghiệp", "Chỉ thấy dx-core-api"),
        ("app-lab", "Miền nghiệp vụ vận hành CNTT", "Chỉ thấy dx-core-api"),
        ("dx-host", "Ghép Core và ứng dụng, tạo gói chạy", "Nạp app theo DXOS_APP"),
    ], widths=[3.2, 8.0, 5.8], font_size=8.5)
    add_doc_para(doc, "Ứng dụng không được phép import entity hoặc implementation bên trong dx-core. Core cũng không import ứng dụng. Các entity dùng chung được che sau các view và facade như CaseView, WorkspaceView, WorkflowOperations, ObjectStorage, UserContext và KnowledgeGateway. Vì vậy, tính độc lập không chỉ là tuyên bố kiến trúc mà được compiler cưỡng chế.")
    doc.add_paragraph("2.2 Cơ chế cắm ứng dụng", style="Heading 2")
    add_doc_para(doc, "Mỗi ứng dụng khai báo vai, loại workflow, cách làm giàu hộp duyệt, loại thông báo, kênh chat, view BI, namespace tri thức, xử lý tệp và chỉ báo sức khỏe qua các SPI. Khi DXOS_APP thay đổi, host chỉ nạp entity, migration, controller và lịch chạy của ứng dụng được chọn. Gateway ghép cấu hình chung với route riêng của app; web đọc manifest từ GET api v1 app để nạp đúng menu và màn hình.")
    add_doc_callout(doc, "BẰNG CHỨNG TÁI DÙNG", "DX Lab được thêm sau DX Farm mà không sửa lõi dùng chung; một ứng dụng mẫu tối thiểu cũng có thể khởi động chỉ với hợp đồng AppModule.", fill=PALE_GREEN, accent=TEAL)

    # Page 4
    add_doc_page_break(doc)
    add_doc_heading(doc, "3 Dịch vụ dùng chung và luồng gọi", "Trang 4")
    add_doc_table(doc, ["Dịch vụ lõi", "Cách cung cấp", "Cách ứng dụng sử dụng"], [
        ("Identity và SSO", "Keycloak realm dxos; OIDC; Core đồng bộ vai do app khai báo", "Farm và Lab dùng vai khác nhau nhưng chung phiên đăng nhập và cơ chế RBAC"),
        ("API Gateway", "APISIX là điểm vào duy nhất, kiểm JWT, định tuyến và giới hạn tốc độ", "Mọi web, mobile và API đi qua cổng 9080; backend kiểm JWT lần hai"),
        ("Dữ liệu", "PostgreSQL và Flyway cho dữ liệu cấu trúc; MinIO cho tệp riêng tư", "App có migration riêng; client nhận link tệp có hạn"),
        ("Workflow", "Case type cấu hình được; workflow_event chỉ thêm; hộp duyệt dùng chung", "Farm có năm loại duyệt; Lab dùng ticket, mượn thiết bị và đặt phòng"),
        ("Thông báo", "Publisher dùng chung, kênh chat do app khai báo", "Lỗi Mattermost không làm hỏng transaction nghiệp vụ"),
        ("AI và tri thức", "Gateway nội bộ, circuit breaker, RAG theo namespace", "Farm và Lab dùng corpus và năng lực Intelligence riêng"),
    ], widths=[3.0, 7.1, 6.9], font_size=8.0)
    doc.add_paragraph("3.1 Luồng yêu cầu DX Farm", style="Heading 2")
    add_doc_para(doc, "Người dùng đăng nhập qua Keycloak, gửi báo cáo bệnh kèm ảnh qua APISIX. Backend xác thực lại token và phạm vi nông hộ, yêu cầu AI kiểm chất lượng và xác nhận ảnh lúa trước khi lưu, xóa metadata ảnh rồi ghi tệp vào MinIO và bản ghi vào PostgreSQL. Khi kỹ sư yêu cầu chẩn đoán, Backend cấp link nội bộ có hạn cho AI, nhận kết quả, ghi Diagnosis và mở workflow case. Kỹ sư duyệt; quy tắc thuốc, liều và PHI được lấy từ dữ liệu có cấu trúc.")
    doc.add_paragraph("3.2 Luồng yêu cầu DX Lab", style="Heading 2")
    add_doc_para(doc, "Sinh viên nhập mô tả sự cố và nhận gợi ý loại, ưu tiên, lý do cùng bài tri thức liên quan. Ticket được lưu với hạn phản hồi và hạn xử lý theo SLA, đồng thời mở case và phát thông báo. Khi hai kỹ thuật viên cùng nhận việc, khóa phiên bản bảo đảm chỉ một người thành công; phía còn lại nhận 409. Timeline ghi nhận từng bước, bình luận nội bộ được ẩn với người yêu cầu, và sinh viên xác nhận kết quả sau khi kỹ thuật viên hoàn tất.")
    add_doc_callout(doc, "NGUYÊN TẮC", "Client chỉ biết Gateway. AI không truy cập trực tiếp cơ sở dữ liệu nghiệp vụ. PostgreSQL là nguồn sự thật; RAG phục vụ tra cứu và diễn đạt, không thay thế dữ kiện có cấu trúc.", fill=PALE, accent=BLUE)

    # Page 5
    add_doc_page_break(doc)
    add_doc_heading(doc, "4 Không gian năng lực H P D I", "Trang 5")
    add_doc_para(doc, "H P D I được thể hiện qua hành vi vận hành, không phải bốn nhãn trên dashboard. Cùng một tập dịch vụ lõi tạo ra năng lực khác nhau tùy ứng dụng, vai trò và dữ liệu.")
    add_doc_table(doc, ["Không gian", "DX Farm", "DX Lab", "Minh chứng kỹ thuật"], [
        ("H Human", "Nông dân chỉ thấy hộ mình; kỹ sư và quản trị theo vai; mobile tách khỏi web", "Sinh viên chỉ thấy ticket của mình; kỹ thuật viên thấy hàng đợi; ghi chú nội bộ được che", "Keycloak, workspace, UserContext, RBAC và policy theo client"),
        ("P Process", "Duyệt thửa, vụ, chẩn đoán, khai phun; khóa thu hoạch theo PHI", "Ticket theo SLA; duyệt mượn thiết bị; đặt phòng chống chồng giờ", "Workflow case và event; state transition; optimistic locking; ràng buộc DB"),
        ("D Data", "View bản ghi đã duyệt, dữ liệu thời tiết có snapshot, ảnh trong object storage", "View SLA, tải kỹ thuật viên, mức dùng phòng; tệp đính kèm", "PostgreSQL, MinIO, Superset, metadata và provenance"),
        ("I Intelligence", "Crop classifier, detector bệnh, RAG và diễn đạt bằng LLM", "Luật phân loại giải thích được và RAG quy chế", "Ngưỡng, lý do, nguồn trích dẫn, human in the loop và fallback"),
    ], widths=[2.0, 4.7, 4.7, 5.6], font_size=7.8)
    doc.add_paragraph("4.1 Process và poka yoke", style="Heading 2")
    add_doc_para(doc, "Các quy tắc an toàn được đặt ở Backend hoặc cơ sở dữ liệu thay vì chỉ ẩn nút trên giao diện. Với Farm, thu hoạch bị chặn khi chưa hết PHI; bệnh chưa có treatment rule được duyệt chỉ ở trạng thái review_required và không có liều hay PHI. Với Lab, đặt phòng trùng giờ bị chặn bằng ràng buộc EXCLUDE USING gist, nên ngay cả hai yêu cầu đồng thời cũng chỉ có một lịch hiệu lực.")
    doc.add_paragraph("4.2 Data và khả năng kiểm chứng", style="Heading 2")
    add_doc_para(doc, "Dữ liệu nghiệp vụ có nguồn gốc rõ: dữ liệu Farm được sinh từ bộ JSON có trạng thái phê duyệt; chỉ record eligible_for_seed mới vào seed. Tài liệu RAG tách namespace farm và lab, mang thông tin nguồn và giấy phép. BI chỉ đọc các view cho phép, không đọc trực tiếp bảng nghiệp vụ. Khi API thời tiết không sẵn sàng, hệ thống trả snapshot và đánh dấu isSnapshot để người dùng biết chất lượng dữ liệu.")
    add_doc_callout(doc, "KHÁC CRUD", "Giá trị nằm ở phân quyền theo phạm vi, workflow có lịch sử, ràng buộc đồng thời, dữ liệu có provenance và Intelligence giải thích được; dashboard chỉ là một bề mặt quan sát.", fill=PALE_GOLD, accent="C28A18")

    # Page 6
    add_doc_page_break(doc)
    add_doc_heading(doc, "5 Ứng dụng nghiệp vụ và kịch bản đầu cuối", "Trang 6")
    doc.add_paragraph("5.1 DX Farm", style="Heading 2")
    add_doc_table(doc, ["Bước", "Hành động", "Dịch vụ lõi tham gia", "Kết quả"], [
        ("1", "Nông dân gửi ảnh báo bệnh", "SSO, Gateway, Storage, Data", "Ảnh được kiểm, làm sạch metadata và lưu riêng tư"),
        ("2", "AI phân tích", "AI Gateway, Knowledge, Storage", "Trả lớp bệnh, độ tin cậy và căn cứ"),
        ("3", "Kỹ sư duyệt", "Workflow, Notification", "Chẩn đoán trở thành quyết định có người chịu trách nhiệm"),
        ("4", "Khai phun thuốc", "Structured Data, Workflow", "Chỉ rule đã duyệt mới cung cấp liều và PHI"),
        ("5", "Thu hoạch", "Poka yoke, Data", "Bị chặn trước hạn; được phép sau khi thỏa điều kiện"),
    ], widths=[1.2, 4.5, 5.0, 6.3], font_size=8.1)
    doc.add_paragraph("5.2 DX Lab", style="Heading 2")
    add_doc_table(doc, ["Bước", "Hành động", "Dịch vụ lõi tham gia", "Kết quả"], [
        ("1", "Sinh viên mô tả sự cố", "Rule Intelligence, Knowledge", "Gợi ý loại, ưu tiên, lý do và bài liên quan"),
        ("2", "Tạo ticket kèm ảnh", "Gateway, Storage, Data, Workflow", "Sinh mã, hạn SLA, case và timeline"),
        ("3", "Kỹ thuật viên nhận và xử lý", "Workflow, RBAC, locking", "Một người nhận thành công; trao đổi công khai và nội bộ"),
        ("4", "Hoàn tất và xác nhận", "Workflow, Notification", "Có ghi chú, nguyên nhân gốc và đánh giá"),
        ("5", "Đặt phòng hoặc mượn thiết bị", "Workflow, DB constraints", "Duyệt theo vai và chặn xung đột"),
    ], widths=[1.2, 4.5, 5.0, 6.3], font_size=8.1)
    doc.add_paragraph("5.3 Chuyển ứng dụng trên cùng nền tảng", style="Heading 2")
    add_doc_para(doc, "Lệnh khởi động với farm hoặc lab thay đổi biến DXOS_APP. Keycloak, PostgreSQL, MinIO và APISIX vẫn là hạ tầng chung; Backend chỉ nạp module ứng dụng được chọn, gateway chỉ có route của app đang chạy, Superset đăng ký đúng dataset và job RAG nạp đúng corpus. Route của ứng dụng còn lại trả 404. Bảng dữ liệu của app trước vẫn tồn tại nhưng không được sử dụng.")
    add_doc_callout(doc, "GIÁ TRỊ MINH HỌA", "Việc chuyển từ bài toán nông nghiệp sang vận hành CNTT bằng cùng Core là minh chứng trực tiếp rằng nền tảng không phụ thuộc một giao diện hay một miền dữ liệu.", fill=PALE_GREEN, accent=TEAL)

    # Page 7
    add_doc_page_break(doc)
    add_doc_heading(doc, "6 Trí tuệ nhân tạo dữ liệu và giới hạn", "Trang 7")
    doc.add_paragraph("6.1 Pipeline Intelligence của DX Farm", style="Heading 2")
    add_doc_table(doc, ["Giai đoạn", "Công nghệ và dữ liệu", "Kiểm soát"], [
        ("Kiểm ảnh", "Kiểm chất lượng và crop classifier YOLO26s cls dạng ONNX", "Rice từ 0,80; chênh lớp khác từ 0,20; lỗi dịch vụ trả 503"),
        ("Nhận diện bệnh", "YOLOv8n RiceLeafBD 3.1.0 trên năm bệnh lá lúa", "Confidence trong code 0,70; ảnh không đạt bị chặn"),
        ("Tra cứu", "pgvector và embedding paraphrase multilingual MiniLM L12 v2", "Corpus Farm gồm tài liệu có nguồn; tách namespace"),
        ("Diễn đạt", "Qwen2.5 1.5B qua Ollama và LangGraph", "Không được sinh số thuốc, liều hoặc PHI"),
        ("Quyết định", "Kỹ sư duyệt chẩn đoán và treatment rule", "Human in the loop; nhãn chốt được lưu cho đánh giá sau"),
    ], widths=[2.8, 8.0, 6.2], font_size=8.1)
    doc.add_paragraph("6.2 Intelligence của DX Lab", style="Heading 2")
    add_doc_para(doc, "Lab sử dụng luật từ khóa để gợi ý loại và ưu tiên. Kết quả luôn kèm từ khóa khớp và lý do, nên có thể giải thích và kiểm tra. RAG tìm bài quy chế hoặc hướng dẫn liên quan, nhưng người dùng vẫn chọn nội dung ticket và kỹ thuật viên vẫn chịu trách nhiệm xử lý. Nếu tri thức không sẵn sàng, luật phân loại vẫn hoạt động và knowledgeAvailable được đặt false.")
    doc.add_paragraph("6.3 Giới hạn và đánh giá trung thực", style="Heading 2")
    add_doc_bullets(doc, [
        "Ngưỡng 0,70 của detector chưa được hiệu chỉnh bằng một tập ảnh thực địa độc lập theo từng lớp.",
        "Crop classifier còn nhận nhầm một số ảnh ngô, cỏ và lúa mì; kiểm soát đầu vào giúp giảm rủi ro nhưng chưa loại bỏ hoàn toàn.",
        "Mô hình chưa bao phủ sâu hại; treatment rule được duyệt mới có cho đạo ôn và khô vằn.",
        "Luật phân loại Lab là luật tay và chưa có chỉ số độ chính xác trên dữ liệu vận hành thật.",
        "LLM có nhánh evidence fallback; quyết định an toàn luôn dựa trên dữ liệu có cấu trúc và người duyệt.",
    ])
    add_doc_callout(doc, "NGUYÊN TẮC AI", "Deterministic first, human in the loop và có căn cứ. AI hỗ trợ nhận diện, tra cứu và ưu tiên; không tự quyết định phun thuốc, liều lượng, PHI hay kết luận nghiệp vụ cuối cùng.", fill=PALE_GOLD, accent="C28A18")

    # Page 8
    add_doc_page_break(doc)
    add_doc_heading(doc, "7 Bảo mật kiểm thử và khả năng phục hồi", "Trang 8")
    doc.add_paragraph("7.1 Bảo mật theo nhiều lớp", style="Heading 2")
    add_doc_table(doc, ["Lớp kiểm soát", "Biện pháp"], [
        ("Biên hệ thống", "APISIX kiểm JWT, định tuyến theo app, giới hạn tốc độ; chỉ Gateway được publish cho client"),
        ("Ứng dụng", "Backend xác thực JWT lần hai, RBAC, phạm vi workspace và policy theo loại client"),
        ("Dữ liệu", "MinIO bucket riêng tư, link có hạn; BI chỉ đọc view; migration bằng Flyway"),
        ("Bí mật", "Biến môi trường, tệp env example, dừng sớm khi thiếu cấu hình; không hard code secret"),
        ("Tính toàn vẹn", "Optimistic locking, workflow event chỉ thêm, constraint chống trùng và lỗi ổn định cho client"),
    ], widths=[4.1, 12.9], font_size=8.5)
    doc.add_paragraph("7.2 Kiểm thử và kiểm chứng", style="Heading 2")
    add_doc_table(doc, ["Phạm vi", "Bằng chứng gần nhất", "Ý nghĩa"], [
        ("Backend", "Runbook ngày 13 10 2026 ghi nhận mvnw clean verify đạt 67 test", "Kiểm hợp đồng, workflow, quyền và tích hợp dữ liệu"),
        ("AI", "Tài liệu kiểm chứng CV ghi nhận 64 test, lint và dependency check đạt", "Kiểm loader, policy, precheck, predict, diagnose và lỗi dịch vụ"),
        ("E2E Farm", "10 trên 10 bước đạt với require rule", "Chứng minh vai, duyệt, idempotency, kiểm ảnh, chẩn đoán và PHI"),
        ("E2E Lab", "16 trên 16 bước đạt với require knowledge", "Chứng minh ticket, SLA, thiết bị, phòng và tri thức"),
        ("Ranh giới app", "Route app không chạy trả 404; app sai dừng sớm; đổi app không lỗi migration", "Chứng minh lựa chọn ứng dụng và tính cô lập khi chạy"),
    ], widths=[3.0, 7.0, 7.0], font_size=8.0)
    doc.add_paragraph("7.3 Khả năng phục hồi", style="Heading 2")
    add_doc_para(doc, "Open Meteo lỗi thì dùng snapshot có đánh dấu; LLM lỗi thì trả evidence fallback; gợi ý Lab vẫn chạy bằng luật; Mattermost gửi bất đồng bộ; mobile giữ hàng đợi và gửi lại; portal công khai trạng thái dịch vụ. Các đường dự phòng ưu tiên duy trì nghiệp vụ cốt lõi và làm rõ chất lượng dữ liệu thay vì che giấu lỗi.")
    add_doc_callout(doc, "LƯU Ý BẰNG CHỨNG", "Một số kiểm tra giao diện và dashboard vẫn là kiểm tay. Sản phẩm chưa có thử tải, kiểm thử xâm nhập hoặc báo cáo đánh giá mô hình độc lập.", fill="FBECEC", accent=RED)

    # Page 9
    add_doc_page_break(doc)
    add_doc_heading(doc, "8 Nguồn mở giấy phép và khả năng tái lập", "Trang 9")
    add_doc_para(doc, "Proof of FOSS được coi là điều kiện nền của sản phẩm. Repository có LICENSE, bảng rà soát dependency và license, SBOM CycloneDX, lockfile, CHANGELOG, hướng dẫn đóng góp và pull request template. Repository có lịch sử commit thực tế, nhưng bản kiểm tra cục bộ chưa có tag; tag và release phải được tạo trước khi nộp. Quy trình dựng từ source dùng Docker Compose, Maven wrapper và dữ liệu demo có thể nạp lại.")
    doc.add_paragraph("8.1 Thành phần và giấy phép chính", style="Heading 2")
    add_doc_table(doc, ["Nhóm", "Thành phần", "Giấy phép hoặc cách phân phối"], [
        ("Mã tự viết", "Core, backend ứng dụng, web, mobile và hạ tầng cấu hình", "Apache 2.0; file Java và Python có SPDX theo quy ước dự án"),
        ("AI", "Ultralytics và các checkpoint CV", "AGPL 3.0; thư mục ai phân phối theo AGPL 3.0"),
        ("Định danh và Gateway", "Keycloak, Apache APISIX", "Apache 2.0"),
        ("Dữ liệu", "PostgreSQL, pgvector, MinIO", "PostgreSQL License; MinIO AGPL 3.0"),
        ("BI và Chat", "Apache Superset, Mattermost Team Edition", "Apache 2.0; Mattermost theo điều kiện nguồn mở được ghi nhận"),
        ("Mô hình và dữ liệu", "Qwen, MiniLM, RiceLeafDiseaseBD và tài liệu RAG", "Ghi phiên bản, nguồn, vai trò và giấy phép trong LICENSES và metadata"),
    ], widths=[3.1, 6.0, 7.9], font_size=8.0)
    doc.add_paragraph("8.2 Dựng và chạy từ source", style="Heading 2")
    add_doc_bullets(doc, [
        "Tạo cấu hình từ deploy .env.example và điền các biến bắt buộc; Compose dừng và chỉ rõ biến còn thiếu.",
        "Chạy script start với farm hoặc lab để dựng stack, chờ health, nạp tri thức và dữ liệu demo.",
        "Dùng reset để xóa dữ liệu triển khai và tái lập môi trường sạch; model Ollama có thể được giữ lại để giảm thời gian.",
        "Swagger sinh từ code mô tả API; tài khoản demo và địa chỉ dịch vụ được ghi trong README.",
        "CI tách kiểm tra backend, AI, web, mobile, hạ tầng và E2E cho cả hai ứng dụng.",
    ])
    doc.add_paragraph("8.3 Phát triển cộng đồng", style="Heading 2")
    add_doc_para(doc, "CONTRIBUTING mô tả quy trình nhánh và pull request, kiểm tra theo module và cách thêm ứng dụng mới. Người đóng góp tạo module chỉ phụ thuộc dx-core-api, cài AppModule và các SPI cần thiết, thêm migration riêng, route gateway, dataset BI, corpus tri thức và giao diện. app-lab là ví dụ hoàn chỉnh để đối chiếu.")
    add_doc_callout(doc, "RỦI RO GIẤY PHÉP", "AGPL được cô lập và công bố rõ. AI Service là tiến trình riêng giao tiếp với Backend qua HTTP; mọi thành phần bên ngoài cần tiếp tục được khóa phiên bản, ghi nguồn và rà soát khi nâng cấp.", fill=PALE_GOLD, accent="C28A18")

    # Page 10
    add_doc_page_break(doc)
    add_doc_heading(doc, "9 Đánh giá giới hạn và hướng phát triển", "Trang 10")
    doc.add_paragraph("9.1 Đối chiếu yêu cầu đề thi", style="Heading 2")
    add_doc_table(doc, ["Yêu cầu", "Mức đáp ứng", "Bằng chứng chính"], [
        ("Core headless và ba tầng", "Đáp ứng", "Ranh giới module, SPI, Gateway, Identity, Data, Workflow và hai app"),
        ("H P D I", "Đáp ứng", "Vai và phạm vi; workflow; PostgreSQL MinIO BI; CV RAG và luật"),
        ("Ứng dụng thực tế", "Đáp ứng", "DX Farm và DX Lab có dữ liệu mẫu và kịch bản đầu cuối"),
        ("API và triển khai", "Đáp ứng", "REST OpenAPI, Docker Compose, script start reset và dữ liệu demo"),
        ("Kiểm thử", "Đáp ứng mức MVP", "Test backend AI và E2E hai app; còn thiếu thử tải và bảo mật chuyên sâu"),
        ("PoF nguồn mở", "Đáp ứng phần lớn", "LICENSE, LICENSES, SBOM, CHANGELOG và CONTRIBUTING; chưa có tag trong bản kiểm tra"),
    ], widths=[4.2, 3.0, 9.8], font_size=8.1)
    doc.add_paragraph("9.2 Giới hạn hiện tại", style="Heading 2")
    add_doc_bullets(doc, [
        "Dashboard Superset của Lab và mobile cho Lab còn trong backlog; một số kiểm thử giao diện vẫn thực hiện thủ công.",
        "AI Service hiện là một tiến trình, chưa có hàng đợi suy luận; Superset dùng SQLite cho metadata; chưa có thử tải.",
        "Các lịch nền cần cơ chế leader hoặc distributed lock khi chạy nhiều bản backend.",
        "Rule thuốc được duyệt còn hẹp; mô hình ảnh cần đánh giá độc lập trên dữ liệu thực địa và báo cáo theo từng lớp.",
        "Snapshot thời tiết chỉ bao phủ một số điểm tham chiếu, nên cần mở rộng phạm vi và quy trình cập nhật provenance.",
    ])
    doc.add_paragraph("9.3 Hướng phát triển", style="Heading 2")
    add_doc_para(doc, "Ưu tiên tiếp theo là hoàn thiện dashboard Lab, bổ sung test tự động cho giao diện, thiết lập đánh giá mô hình và data card, đưa suy luận AI sang hàng đợi worker, bổ sung distributed scheduling, thử tải và kiểm thử bảo mật. Với cộng đồng, cần công bố release có checksum, duy trì SBOM theo mỗi phiên bản và dùng issue tracker để quản lý backlog theo module.")
    add_doc_callout(doc, "KẾT LUẬN", "DX OS thể hiện đúng tư duy nền tảng: Core cung cấp dịch vụ dùng chung và giữ nguyên khi đổi ứng dụng; H P D I liên kết con người, quy trình, dữ liệu và trí tuệ; hai ứng dụng chứng minh giá trị sử dụng và khả năng mở rộng. Các giới hạn được công bố để Hội đồng phân biệt rõ phần đã kiểm chứng với hướng phát triển.", fill=PALE_GREEN, accent=TEAL)
    doc.add_paragraph("Tài liệu tham chiếu nội bộ", style="Heading 2")
    refs = [
        "Đề thi PMNM HaUI 2026 bản chính thức, mục 3 đến 11.",
        "README dự án; ARCHITECTURE_OPEN_CORE; REPORT_OUTLINE; DEMO_SCRIPT; COUNCIL_QA.",
        "LICENSES; CONTRIBUTING; README của backend, app-farm, app-lab và ai.",
        "RUNBOOK_KIEM_THU và CV_RICE_IMAGE_VALIDATION_SUMMARY, cập nhật đến ngày 13 10 2026.",
    ]
    add_doc_bullets(doc, refs)

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc.save(DOCX_OUT)
    return DOCX_OUT


# PDF helpers
FONT_REG = "DXArial"
FONT_BOLD = "DXArialBold"


def register_fonts():
    arial = Path(r"C:\Windows\Fonts\arial.ttf")
    arial_bold = Path(r"C:\Windows\Fonts\arialbd.ttf")
    if not arial.exists() or not arial_bold.exists():
        raise FileNotFoundError("Arial fonts are required for Vietnamese PDF output")
    pdfmetrics.registerFont(TTFont(FONT_REG, str(arial)))
    pdfmetrics.registerFont(TTFont(FONT_BOLD, str(arial_bold)))


def pdf_para(text, style):
    return Paragraph(escape(text).replace("\n", "<br/>"), style)


def pdf_table(headers, rows, widths, styles, font_size=7.7):
    data = [[Paragraph(f"<b>{escape(str(h))}</b>", styles["table_head"]) for h in headers]]
    for row in rows:
        data.append([Paragraph(escape(str(v)), styles["table_cell"]) for v in row])
    t = Table(data, colWidths=[w * mm for w in widths], repeatRows=1, hAlign="CENTER")
    ts = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + NAVY)),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), FONT_BOLD),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#CCD5DB")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]
    for i in range(1, len(data)):
        if i % 2 == 0:
            ts.append(("BACKGROUND", (0, i), (-1, i), colors.HexColor("#" + LIGHT_GRAY)))
    t.setStyle(TableStyle(ts))
    return t


def pdf_callout(label, text, styles, fill=PALE, accent=BLUE):
    data = [[Paragraph(f"<b>{escape(label)}</b>", styles["callout_label"]), Paragraph(escape(text), styles["callout_text"])]]
    t = Table(data, colWidths=[31 * mm, 139 * mm], hAlign="CENTER")
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (0, 0), colors.HexColor("#" + accent)),
        ("BACKGROUND", (1, 0), (1, 0), colors.HexColor("#" + fill)),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 7),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
    ]))
    return t


def pdf_bullets(items, styles):
    flows = []
    for item in items:
        flows.append(Paragraph("• " + escape(item), styles["bullet"]))
    return flows


def on_pdf_page(canvas, doc):
    canvas.saveState()
    width, height = A4
    canvas.setFont(FONT_BOLD, 7.2)
    canvas.setFillColor(colors.HexColor("#" + BLUE))
    canvas.drawRightString(width - 18 * mm, height - 9.5 * mm, "DX OS OPEN CORE   |   BÁO CÁO KỸ THUẬT")
    canvas.setFont(FONT_REG, 8)
    canvas.setFillColor(colors.HexColor("#" + GRAY))
    canvas.drawRightString(width - 18 * mm, 9 * mm, f"Trang {doc.page}")
    canvas.restoreState()


def build_pdf():
    register_fonts()
    ss = getSampleStyleSheet()
    styles = {
        "title": ParagraphStyle("TitleDX", parent=ss["Title"], fontName=FONT_BOLD, fontSize=24, leading=28, textColor=colors.HexColor("#" + NAVY), alignment=TA_CENTER, spaceAfter=8),
        "subtitle": ParagraphStyle("SubtitleDX", parent=ss["Normal"], fontName=FONT_REG, fontSize=11, leading=14, textColor=colors.HexColor("#" + GRAY), alignment=TA_CENTER, spaceAfter=12),
        "h1": ParagraphStyle("H1DX", parent=ss["Heading1"], fontName=FONT_BOLD, fontSize=16, leading=19, textColor=colors.HexColor("#" + NAVY), spaceAfter=6),
        "h2": ParagraphStyle("H2DX", parent=ss["Heading2"], fontName=FONT_BOLD, fontSize=11, leading=13, textColor=colors.HexColor("#" + BLUE), spaceBefore=5, spaceAfter=3),
        "body": ParagraphStyle("BodyDX", parent=ss["BodyText"], fontName=FONT_REG, fontSize=8.9, leading=11.3, textColor=colors.HexColor("#24313A"), alignment=TA_JUSTIFY, spaceAfter=4),
        "bullet": ParagraphStyle("BulletDX", parent=ss["BodyText"], fontName=FONT_REG, fontSize=8.5, leading=10.4, leftIndent=10, firstLineIndent=-7, textColor=colors.HexColor("#24313A"), spaceAfter=2.2),
        "table_head": ParagraphStyle("THDX", parent=ss["BodyText"], fontName=FONT_BOLD, fontSize=7.2, leading=8.4, textColor=colors.white, alignment=TA_CENTER),
        "table_cell": ParagraphStyle("TCDX", parent=ss["BodyText"], fontName=FONT_REG, fontSize=7.35, leading=8.7, textColor=colors.HexColor("#24313A"), alignment=TA_LEFT),
        "callout_label": ParagraphStyle("CLDX", parent=ss["BodyText"], fontName=FONT_BOLD, fontSize=8, leading=9.5, textColor=colors.white, alignment=TA_CENTER),
        "callout_text": ParagraphStyle("CTDX", parent=ss["BodyText"], fontName=FONT_REG, fontSize=8.2, leading=10, textColor=colors.HexColor("#" + NAVY), alignment=TA_LEFT),
        "kicker": ParagraphStyle("KickerDX", parent=ss["BodyText"], fontName=FONT_BOLD, fontSize=7.8, leading=9, textColor=colors.HexColor("#" + TEAL), spaceAfter=2),
    }

    frame = Frame(18 * mm, 14 * mm, A4[0] - 36 * mm, A4[1] - 29 * mm, leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0)
    doc = BaseDocTemplate(str(PDF_OUT), pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=15 * mm, bottomMargin=14 * mm, title="Báo cáo kỹ thuật DX OS Open Core cho DX Farm và DX Lab", author="Nhóm phát triển DX OS")
    doc.addPageTemplates(PageTemplate(id="main", frames=[frame], onPage=on_pdf_page))
    story = []

    def h1(text, kicker=None):
        if kicker:
            story.append(Paragraph(escape(kicker.upper()), styles["kicker"]))
        story.append(Paragraph(escape(text), styles["h1"]))

    def h2(text): story.append(Paragraph(escape(text), styles["h2"]))
    def body(text): story.append(pdf_para(text, styles["body"]))
    def table(headers, rows, widths):
        story.append(pdf_table(headers, rows, widths, styles))
        story.append(Spacer(1, 3))
    def callout(label, text, fill=PALE, accent=BLUE):
        story.append(pdf_callout(label, text, styles, fill, accent))
        story.append(Spacer(1, 4))
    def bullets(items): story.extend(pdf_bullets(items, styles))
    def br(): story.append(PageBreak())

    # Page 1
    story.append(Spacer(1, 14 * mm))
    story.append(Paragraph("PMNM HaUI 2026", styles["kicker"]))
    story.append(Paragraph("Báo cáo kỹ thuật DX OS Open Core cho DX Farm và DX Lab", styles["title"]))
    story.append(Paragraph("Nền tảng lõi dùng chung cho chuyển đổi số nông nghiệp và vận hành CNTT đại học", styles["subtitle"]))
    callout("KẾT LUẬN CHÍNH", "Sản phẩm đáp ứng định hướng Open Core trước và ứng dụng sau: Tầng 1 cung cấp dịch vụ headless dùng chung; Tầng 2 tổ chức năng lực H P D I; Tầng 3 chứng minh khả năng tái sử dụng bằng hai ứng dụng DX Farm và DX Lab trên cùng một lõi.", PALE_GREEN, TEAL)
    h2("Tóm tắt điều hành")
    body("DX OS được tổ chức như một nền tảng tích hợp thay vì một website nghiệp vụ đơn lẻ. Lõi dùng chung gồm định danh và SSO, cổng API, dữ liệu cấu trúc và phi cấu trúc, workflow, thông báo, lưu trữ, cổng AI và tri thức. Ứng dụng chỉ phụ thuộc hợp đồng trong dx-core-api; việc chọn ứng dụng thực thi qua biến DXOS_APP. Ranh giới này biến tính độc lập của Core thành một thuộc tính có thể kiểm tra khi biên dịch và khi chạy.")
    body("DX Farm minh họa chuỗi báo bệnh bằng ảnh, chẩn đoán có người duyệt, khai phun và khóa thu hoạch theo thời gian cách ly. DX Lab minh họa hỗ trợ CNTT theo SLA, mượn thiết bị, đặt phòng máy, tri thức và thống kê. Hai miền nghiệp vụ khác nhau nhưng dùng cùng dịch vụ nền, qua đó chứng minh Core có vai trò thật và không bị khóa vào một giao diện cụ thể.")
    table(["Trọng tâm đề thi", "Minh chứng trong sản phẩm"], [
        ("Open Core headless", "dx-core, Keycloak, APISIX, PostgreSQL, MinIO và các facade dùng chung"),
        ("Kiến trúc 3 tầng", "Core dùng chung, không gian H P D I, hai ứng dụng nghiệp vụ"),
        ("PoF nguồn mở", "License, bảng dependency, SBOM và quy trình dựng từ source; cần tạo tag và release nộp bài"),
        ("Intelligence có căn cứ", "CV và RAG ở Farm; luật giải thích được và RAG ở Lab; con người chốt quyết định"),
    ], [47, 123])
    story.append(Paragraph("Ngày hoàn thiện báo cáo 13 tháng 10 năm 2026", styles["subtitle"]))

    # Page 2
    br(); h1("1 Bài toán phạm vi và tiêu chí thành công", "Trang 2")
    body("Đề thi yêu cầu xây dựng DX Lab HaUI theo tư duy Open Core trước, ứng dụng sau. Vì vậy, phạm vi sản phẩm không dừng ở một dashboard hoặc ứng dụng CRUD. Nền tảng phải cung cấp các dịch vụ dùng chung cho lớp H P D I và cho ít nhất một ứng dụng nghiệp vụ, đồng thời chứng minh được khả năng cài đặt, vận hành và phát triển bằng công nghệ nguồn mở.")
    h2("1.1 Mục tiêu thiết kế")
    bullets(["Tách phần lõi headless khỏi giao diện và khỏi mô hình dữ liệu đặc thù của từng ứng dụng.", "Cung cấp một cửa truy cập qua API Gateway, xác thực tập trung, phân quyền theo vai và phạm vi dữ liệu.", "Mô hình hóa quy trình dưới dạng case và sự kiện để lưu được trạng thái, người xử lý, thời gian và lịch sử thay đổi.", "Dùng PostgreSQL làm nguồn sự thật; object storage quản lý tệp; BI chỉ đọc các view đã kiểm soát.", "Đặt AI vào quy trình có căn cứ, có ngưỡng, có đường dự phòng và có con người chịu trách nhiệm quyết định."])
    h2("1.2 Hai ứng dụng minh họa")
    table(["Ứng dụng", "Người dùng chính", "Luồng xương sống", "Giá trị chứng minh"], [("DX Farm", "Nông dân, kỹ sư, quản trị", "Ảnh bệnh lúa → AI → kỹ sư duyệt → khai phun → khóa PHI", "Core phục vụ mobile, web, AI, BI và kiểm soát an toàn nông nghiệp"), ("DX Lab", "Sinh viên, giảng viên, kỹ thuật viên, quản trị", "Gợi ý → tạo ticket → xử lý SLA → xác nhận; mượn thiết bị; đặt phòng", "Cùng Core phục vụ một miền nghiệp vụ đại học hoàn toàn khác")], [23, 36, 61, 50])
    h2("1.3 Tiêu chí thành công")
    body("Thành công được đánh giá trên hai trục. Trục PoF yêu cầu repository công khai, lịch sử phát triển, giấy phép, release, dependency có nguồn gốc và quy trình dựng lại. Trục sản phẩm yêu cầu đúng kiến trúc, chạy ổn định, có giá trị sử dụng, có khả năng mở rộng và trình bày thuyết phục. Báo cáo này ưu tiên các minh chứng có thể kiểm tra thay vì mô tả giao diện.")
    callout("PHẠM VI", "Hai ứng dụng được trình bày để chứng minh khả năng tái sử dụng. Trọng tâm chấm vẫn là nền tảng Open Core và cách ứng dụng khai thác các dịch vụ của lõi.", PALE_GOLD, "C28A18")

    # Page 3
    br(); h1("2 Kiến trúc Open Core ba tầng", "Trang 3")
    body("Kiến trúc đi từ dịch vụ lõi ở dưới lên năng lực số và ứng dụng nghiệp vụ. Luồng phụ thuộc luôn hướng về hợp đồng của Core; giao diện không nằm trong Tầng 1. Cách tổ chức này đáp ứng trực tiếp yêu cầu phân biệt nền tảng dùng chung với sản phẩm giao diện.")
    table(["Tầng", "Vai trò", "Thành phần tiêu biểu", "Minh chứng"], [("Tầng 3 Ứng dụng", "Giải quyết bài toán cụ thể", "DX Farm và DX Lab", "Hai module nghiệp vụ, hai giao diện, cùng một Core"), ("Tầng 2 H P D I", "Biến dịch vụ nền thành năng lực vận hành", "Human, Process, Data, Intelligence", "Vai và workspace; workflow; BI; CV, RAG và luật"), ("Tầng 1 Open Core", "Dịch vụ headless và PaaS cho ứng dụng", "Identity, Gateway, Data, Workflow, Storage, Notification, AI Gateway", "API và SPI dùng chung; không có UI người dùng cuối")], [30, 42, 53, 45])
    h2("2.1 Ranh giới module")
    table(["Module", "Trách nhiệm", "Quan hệ phụ thuộc"], [("dx-core-api", "Interface, DTO, lỗi và hợp đồng mở rộng", "Không phụ thuộc ứng dụng"), ("dx-core", "Cài đặt dịch vụ dùng chung", "Chỉ phụ thuộc dx-core-api"), ("app-farm", "Miền nghiệp vụ nông nghiệp", "Chỉ thấy dx-core-api"), ("app-lab", "Miền nghiệp vụ vận hành CNTT", "Chỉ thấy dx-core-api"), ("dx-host", "Ghép Core và ứng dụng, tạo gói chạy", "Nạp app theo DXOS_APP")], [32, 80, 58])
    body("Ứng dụng không được phép import entity hoặc implementation bên trong dx-core. Core cũng không import ứng dụng. Các entity dùng chung được che sau các view và facade như CaseView, WorkspaceView, WorkflowOperations, ObjectStorage, UserContext và KnowledgeGateway. Vì vậy, tính độc lập không chỉ là tuyên bố kiến trúc mà được compiler cưỡng chế.")
    h2("2.2 Cơ chế cắm ứng dụng")
    body("Mỗi ứng dụng khai báo vai, loại workflow, cách làm giàu hộp duyệt, loại thông báo, kênh chat, view BI, namespace tri thức, xử lý tệp và chỉ báo sức khỏe qua các SPI. Khi DXOS_APP thay đổi, host chỉ nạp entity, migration, controller và lịch chạy của ứng dụng được chọn. Gateway ghép cấu hình chung với route riêng của app; web đọc manifest từ GET api v1 app để nạp đúng menu và màn hình.")
    callout("BẰNG CHỨNG TÁI DÙNG", "DX Lab được thêm sau DX Farm mà không sửa lõi dùng chung; một ứng dụng mẫu tối thiểu cũng có thể khởi động chỉ với hợp đồng AppModule.", PALE_GREEN, TEAL)

    # Page 4
    br(); h1("3 Dịch vụ dùng chung và luồng gọi", "Trang 4")
    table(["Dịch vụ lõi", "Cách cung cấp", "Cách ứng dụng sử dụng"], [("Identity và SSO", "Keycloak realm dxos; OIDC; Core đồng bộ vai do app khai báo", "Farm và Lab dùng vai khác nhau nhưng chung phiên đăng nhập và cơ chế RBAC"), ("API Gateway", "APISIX là điểm vào duy nhất, kiểm JWT, định tuyến và giới hạn tốc độ", "Mọi web, mobile và API đi qua cổng 9080; backend kiểm JWT lần hai"), ("Dữ liệu", "PostgreSQL và Flyway cho dữ liệu cấu trúc; MinIO cho tệp riêng tư", "App có migration riêng; client nhận link tệp có hạn"), ("Workflow", "Case type cấu hình được; workflow_event chỉ thêm; hộp duyệt dùng chung", "Farm có năm loại duyệt; Lab dùng ticket, mượn thiết bị và đặt phòng"), ("Thông báo", "Publisher dùng chung, kênh chat do app khai báo", "Lỗi Mattermost không làm hỏng transaction nghiệp vụ"), ("AI và tri thức", "Gateway nội bộ, circuit breaker, RAG theo namespace", "Farm và Lab dùng corpus và năng lực Intelligence riêng")], [30, 71, 69])
    h2("3.1 Luồng yêu cầu DX Farm")
    body("Người dùng đăng nhập qua Keycloak, gửi báo cáo bệnh kèm ảnh qua APISIX. Backend xác thực lại token và phạm vi nông hộ, yêu cầu AI kiểm chất lượng và xác nhận ảnh lúa trước khi lưu, xóa metadata ảnh rồi ghi tệp vào MinIO và bản ghi vào PostgreSQL. Khi kỹ sư yêu cầu chẩn đoán, Backend cấp link nội bộ có hạn cho AI, nhận kết quả, ghi Diagnosis và mở workflow case. Kỹ sư duyệt; quy tắc thuốc, liều và PHI được lấy từ dữ liệu có cấu trúc.")
    h2("3.2 Luồng yêu cầu DX Lab")
    body("Sinh viên nhập mô tả sự cố và nhận gợi ý loại, ưu tiên, lý do cùng bài tri thức liên quan. Ticket được lưu với hạn phản hồi và hạn xử lý theo SLA, đồng thời mở case và phát thông báo. Khi hai kỹ thuật viên cùng nhận việc, khóa phiên bản bảo đảm chỉ một người thành công; phía còn lại nhận 409. Timeline ghi nhận từng bước, bình luận nội bộ được ẩn với người yêu cầu, và sinh viên xác nhận kết quả sau khi kỹ thuật viên hoàn tất.")
    callout("NGUYÊN TẮC", "Client chỉ biết Gateway. AI không truy cập trực tiếp cơ sở dữ liệu nghiệp vụ. PostgreSQL là nguồn sự thật; RAG phục vụ tra cứu và diễn đạt, không thay thế dữ kiện có cấu trúc.")

    # Page 5
    br(); h1("4 Không gian năng lực H P D I", "Trang 5")
    body("H P D I được thể hiện qua hành vi vận hành, không phải bốn nhãn trên dashboard. Cùng một tập dịch vụ lõi tạo ra năng lực khác nhau tùy ứng dụng, vai trò và dữ liệu.")
    table(["Không gian", "DX Farm", "DX Lab", "Minh chứng kỹ thuật"], [("H Human", "Nông dân chỉ thấy hộ mình; kỹ sư và quản trị theo vai; mobile tách khỏi web", "Sinh viên chỉ thấy ticket của mình; kỹ thuật viên thấy hàng đợi; ghi chú nội bộ được che", "Keycloak, workspace, UserContext, RBAC và policy theo client"), ("P Process", "Duyệt thửa, vụ, chẩn đoán, khai phun; khóa thu hoạch theo PHI", "Ticket theo SLA; duyệt mượn thiết bị; đặt phòng chống chồng giờ", "Workflow case và event; state transition; optimistic locking; ràng buộc DB"), ("D Data", "View bản ghi đã duyệt, dữ liệu thời tiết có snapshot, ảnh trong object storage", "View SLA, tải kỹ thuật viên, mức dùng phòng; tệp đính kèm", "PostgreSQL, MinIO, Superset, metadata và provenance"), ("I Intelligence", "Crop classifier, detector bệnh, RAG và diễn đạt bằng LLM", "Luật phân loại giải thích được và RAG quy chế", "Ngưỡng, lý do, nguồn trích dẫn, human in the loop và fallback")], [20, 47, 47, 56])
    h2("4.1 Process và poka yoke")
    body("Các quy tắc an toàn được đặt ở Backend hoặc cơ sở dữ liệu thay vì chỉ ẩn nút trên giao diện. Với Farm, thu hoạch bị chặn khi chưa hết PHI; bệnh chưa có treatment rule được duyệt chỉ ở trạng thái review_required và không có liều hay PHI. Với Lab, đặt phòng trùng giờ bị chặn bằng ràng buộc EXCLUDE USING gist, nên ngay cả hai yêu cầu đồng thời cũng chỉ có một lịch hiệu lực.")
    h2("4.2 Data và khả năng kiểm chứng")
    body("Dữ liệu nghiệp vụ có nguồn gốc rõ: dữ liệu Farm được sinh từ bộ JSON có trạng thái phê duyệt; chỉ record eligible_for_seed mới vào seed. Tài liệu RAG tách namespace farm và lab, mang thông tin nguồn và giấy phép. BI chỉ đọc các view cho phép, không đọc trực tiếp bảng nghiệp vụ. Khi API thời tiết không sẵn sàng, hệ thống trả snapshot và đánh dấu isSnapshot để người dùng biết chất lượng dữ liệu.")
    callout("KHÁC CRUD", "Giá trị nằm ở phân quyền theo phạm vi, workflow có lịch sử, ràng buộc đồng thời, dữ liệu có provenance và Intelligence giải thích được; dashboard chỉ là một bề mặt quan sát.", PALE_GOLD, "C28A18")

    # Page 6
    br(); h1("5 Ứng dụng nghiệp vụ và kịch bản đầu cuối", "Trang 6")
    h2("5.1 DX Farm")
    table(["Bước", "Hành động", "Dịch vụ lõi tham gia", "Kết quả"], [("1", "Nông dân gửi ảnh báo bệnh", "SSO, Gateway, Storage, Data", "Ảnh được kiểm, làm sạch metadata và lưu riêng tư"), ("2", "AI phân tích", "AI Gateway, Knowledge, Storage", "Trả lớp bệnh, độ tin cậy và căn cứ"), ("3", "Kỹ sư duyệt", "Workflow, Notification", "Chẩn đoán trở thành quyết định có người chịu trách nhiệm"), ("4", "Khai phun thuốc", "Structured Data, Workflow", "Chỉ rule đã duyệt mới cung cấp liều và PHI"), ("5", "Thu hoạch", "Poka yoke, Data", "Bị chặn trước hạn; được phép sau khi thỏa điều kiện")], [12, 45, 50, 63])
    h2("5.2 DX Lab")
    table(["Bước", "Hành động", "Dịch vụ lõi tham gia", "Kết quả"], [("1", "Sinh viên mô tả sự cố", "Rule Intelligence, Knowledge", "Gợi ý loại, ưu tiên, lý do và bài liên quan"), ("2", "Tạo ticket kèm ảnh", "Gateway, Storage, Data, Workflow", "Sinh mã, hạn SLA, case và timeline"), ("3", "Kỹ thuật viên nhận và xử lý", "Workflow, RBAC, locking", "Một người nhận thành công; trao đổi công khai và nội bộ"), ("4", "Hoàn tất và xác nhận", "Workflow, Notification", "Có ghi chú, nguyên nhân gốc và đánh giá"), ("5", "Đặt phòng hoặc mượn thiết bị", "Workflow, DB constraints", "Duyệt theo vai và chặn xung đột")], [12, 45, 50, 63])
    h2("5.3 Chuyển ứng dụng trên cùng nền tảng")
    body("Lệnh khởi động với farm hoặc lab thay đổi biến DXOS_APP. Keycloak, PostgreSQL, MinIO và APISIX vẫn là hạ tầng chung; Backend chỉ nạp module ứng dụng được chọn, gateway chỉ có route của app đang chạy, Superset đăng ký đúng dataset và job RAG nạp đúng corpus. Route của ứng dụng còn lại trả 404. Bảng dữ liệu của app trước vẫn tồn tại nhưng không được sử dụng.")
    callout("GIÁ TRỊ MINH HỌA", "Việc chuyển từ bài toán nông nghiệp sang vận hành CNTT bằng cùng Core là minh chứng trực tiếp rằng nền tảng không phụ thuộc một giao diện hay một miền dữ liệu.", PALE_GREEN, TEAL)

    # Page 7
    br(); h1("6 Trí tuệ nhân tạo dữ liệu và giới hạn", "Trang 7")
    h2("6.1 Pipeline Intelligence của DX Farm")
    table(["Giai đoạn", "Công nghệ và dữ liệu", "Kiểm soát"], [("Kiểm ảnh", "Kiểm chất lượng và crop classifier YOLO26s cls dạng ONNX", "Rice từ 0,80; chênh lớp khác từ 0,20; lỗi dịch vụ trả 503"), ("Nhận diện bệnh", "YOLOv8n RiceLeafBD 3.1.0 trên năm bệnh lá lúa", "Confidence trong code 0,70; ảnh không đạt bị chặn"), ("Tra cứu", "pgvector và embedding paraphrase multilingual MiniLM L12 v2", "Corpus Farm gồm tài liệu có nguồn; tách namespace"), ("Diễn đạt", "Qwen2.5 1.5B qua Ollama và LangGraph", "Không được sinh số thuốc, liều hoặc PHI"), ("Quyết định", "Kỹ sư duyệt chẩn đoán và treatment rule", "Human in the loop; nhãn chốt được lưu cho đánh giá sau")], [28, 80, 62])
    h2("6.2 Intelligence của DX Lab")
    body("Lab sử dụng luật từ khóa để gợi ý loại và ưu tiên. Kết quả luôn kèm từ khóa khớp và lý do, nên có thể giải thích và kiểm tra. RAG tìm bài quy chế hoặc hướng dẫn liên quan, nhưng người dùng vẫn chọn nội dung ticket và kỹ thuật viên vẫn chịu trách nhiệm xử lý. Nếu tri thức không sẵn sàng, luật phân loại vẫn hoạt động và knowledgeAvailable được đặt false.")
    h2("6.3 Giới hạn và đánh giá trung thực")
    bullets(["Ngưỡng 0,70 của detector chưa được hiệu chỉnh bằng một tập ảnh thực địa độc lập theo từng lớp.", "Crop classifier còn nhận nhầm một số ảnh ngô, cỏ và lúa mì; kiểm soát đầu vào giúp giảm rủi ro nhưng chưa loại bỏ hoàn toàn.", "Mô hình chưa bao phủ sâu hại; treatment rule được duyệt mới có cho đạo ôn và khô vằn.", "Luật phân loại Lab là luật tay và chưa có chỉ số độ chính xác trên dữ liệu vận hành thật.", "LLM có nhánh evidence fallback; quyết định an toàn luôn dựa trên dữ liệu có cấu trúc và người duyệt."])
    callout("NGUYÊN TẮC AI", "Deterministic first, human in the loop và có căn cứ. AI hỗ trợ nhận diện, tra cứu và ưu tiên; không tự quyết định phun thuốc, liều lượng, PHI hay kết luận nghiệp vụ cuối cùng.", PALE_GOLD, "C28A18")

    # Page 8
    br(); h1("7 Bảo mật kiểm thử và khả năng phục hồi", "Trang 8")
    h2("7.1 Bảo mật theo nhiều lớp")
    table(["Lớp kiểm soát", "Biện pháp"], [("Biên hệ thống", "APISIX kiểm JWT, định tuyến theo app, giới hạn tốc độ; chỉ Gateway được publish cho client"), ("Ứng dụng", "Backend xác thực JWT lần hai, RBAC, phạm vi workspace và policy theo loại client"), ("Dữ liệu", "MinIO bucket riêng tư, link có hạn; BI chỉ đọc view; migration bằng Flyway"), ("Bí mật", "Biến môi trường, tệp env example, dừng sớm khi thiếu cấu hình; không hard code secret"), ("Tính toàn vẹn", "Optimistic locking, workflow event chỉ thêm, constraint chống trùng và lỗi ổn định cho client")], [41, 129])
    h2("7.2 Kiểm thử và kiểm chứng")
    table(["Phạm vi", "Bằng chứng gần nhất", "Ý nghĩa"], [("Backend", "Runbook ngày 13 10 2026 ghi nhận mvnw clean verify đạt 67 test", "Kiểm hợp đồng, workflow, quyền và tích hợp dữ liệu"), ("AI", "Tài liệu kiểm chứng CV ghi nhận 64 test, lint và dependency check đạt", "Kiểm loader, policy, precheck, predict, diagnose và lỗi dịch vụ"), ("E2E Farm", "10 trên 10 bước đạt với require rule", "Chứng minh vai, duyệt, idempotency, kiểm ảnh, chẩn đoán và PHI"), ("E2E Lab", "16 trên 16 bước đạt với require knowledge", "Chứng minh ticket, SLA, thiết bị, phòng và tri thức"), ("Ranh giới app", "Route app không chạy trả 404; app sai dừng sớm; đổi app không lỗi migration", "Chứng minh lựa chọn ứng dụng và tính cô lập khi chạy")], [30, 70, 70])
    h2("7.3 Khả năng phục hồi")
    body("Open Meteo lỗi thì dùng snapshot có đánh dấu; LLM lỗi thì trả evidence fallback; gợi ý Lab vẫn chạy bằng luật; Mattermost gửi bất đồng bộ; mobile giữ hàng đợi và gửi lại; portal công khai trạng thái dịch vụ. Các đường dự phòng ưu tiên duy trì nghiệp vụ cốt lõi và làm rõ chất lượng dữ liệu thay vì che giấu lỗi.")
    callout("LƯU Ý BẰNG CHỨNG", "Một số kiểm tra giao diện và dashboard vẫn là kiểm tay. Sản phẩm chưa có thử tải, kiểm thử xâm nhập hoặc báo cáo đánh giá mô hình độc lập.", "FBECEC", RED)

    # Page 9
    br(); h1("8 Nguồn mở giấy phép và khả năng tái lập", "Trang 9")
    body("Proof of FOSS được coi là điều kiện nền của sản phẩm. Repository có LICENSE, bảng rà soát dependency và license, SBOM CycloneDX, lockfile, CHANGELOG, hướng dẫn đóng góp và pull request template. Repository có lịch sử commit thực tế, nhưng bản kiểm tra cục bộ chưa có tag; tag và release phải được tạo trước khi nộp. Quy trình dựng từ source dùng Docker Compose, Maven wrapper và dữ liệu demo có thể nạp lại.")
    h2("8.1 Thành phần và giấy phép chính")
    table(["Nhóm", "Thành phần", "Giấy phép hoặc cách phân phối"], [("Mã tự viết", "Core, backend ứng dụng, web, mobile và hạ tầng cấu hình", "Apache 2.0; file Java và Python có SPDX theo quy ước dự án"), ("AI", "Ultralytics và các checkpoint CV", "AGPL 3.0; thư mục ai phân phối theo AGPL 3.0"), ("Định danh và Gateway", "Keycloak, Apache APISIX", "Apache 2.0"), ("Dữ liệu", "PostgreSQL, pgvector, MinIO", "PostgreSQL License; MinIO AGPL 3.0"), ("BI và Chat", "Apache Superset, Mattermost Team Edition", "Apache 2.0; Mattermost theo điều kiện nguồn mở được ghi nhận"), ("Mô hình và dữ liệu", "Qwen, MiniLM, RiceLeafDiseaseBD và tài liệu RAG", "Ghi phiên bản, nguồn, vai trò và giấy phép trong LICENSES và metadata")], [31, 60, 79])
    h2("8.2 Dựng và chạy từ source")
    bullets(["Tạo cấu hình từ deploy .env.example và điền các biến bắt buộc; Compose dừng và chỉ rõ biến còn thiếu.", "Chạy script start với farm hoặc lab để dựng stack, chờ health, nạp tri thức và dữ liệu demo.", "Dùng reset để xóa dữ liệu triển khai và tái lập môi trường sạch; model Ollama có thể được giữ lại để giảm thời gian.", "Swagger sinh từ code mô tả API; tài khoản demo và địa chỉ dịch vụ được ghi trong README.", "CI tách kiểm tra backend, AI, web, mobile, hạ tầng và E2E cho cả hai ứng dụng."])
    h2("8.3 Phát triển cộng đồng")
    body("CONTRIBUTING mô tả quy trình nhánh và pull request, kiểm tra theo module và cách thêm ứng dụng mới. Người đóng góp tạo module chỉ phụ thuộc dx-core-api, cài AppModule và các SPI cần thiết, thêm migration riêng, route gateway, dataset BI, corpus tri thức và giao diện. app-lab là ví dụ hoàn chỉnh để đối chiếu.")
    callout("RỦI RO GIẤY PHÉP", "AGPL được cô lập và công bố rõ. AI Service là tiến trình riêng giao tiếp với Backend qua HTTP; mọi thành phần bên ngoài cần tiếp tục được khóa phiên bản, ghi nguồn và rà soát khi nâng cấp.", PALE_GOLD, "C28A18")

    # Page 10
    br(); h1("9 Đánh giá giới hạn và hướng phát triển", "Trang 10")
    h2("9.1 Đối chiếu yêu cầu đề thi")
    table(["Yêu cầu", "Mức đáp ứng", "Bằng chứng chính"], [("Core headless và ba tầng", "Đáp ứng", "Ranh giới module, SPI, Gateway, Identity, Data, Workflow và hai app"), ("H P D I", "Đáp ứng", "Vai và phạm vi; workflow; PostgreSQL MinIO BI; CV RAG và luật"), ("Ứng dụng thực tế", "Đáp ứng", "DX Farm và DX Lab có dữ liệu mẫu và kịch bản đầu cuối"), ("API và triển khai", "Đáp ứng", "REST OpenAPI, Docker Compose, script start reset và dữ liệu demo"), ("Kiểm thử", "Đáp ứng mức MVP", "Test backend AI và E2E hai app; còn thiếu thử tải và bảo mật chuyên sâu"), ("PoF nguồn mở", "Đáp ứng phần lớn", "LICENSE, LICENSES, SBOM, CHANGELOG và CONTRIBUTING; chưa có tag trong bản kiểm tra")], [42, 30, 98])
    h2("9.2 Giới hạn hiện tại")
    bullets(["Dashboard Superset của Lab và mobile cho Lab còn trong backlog; một số kiểm thử giao diện vẫn thực hiện thủ công.", "AI Service hiện là một tiến trình, chưa có hàng đợi suy luận; Superset dùng SQLite cho metadata; chưa có thử tải.", "Các lịch nền cần cơ chế leader hoặc distributed lock khi chạy nhiều bản backend.", "Rule thuốc được duyệt còn hẹp; mô hình ảnh cần đánh giá độc lập trên dữ liệu thực địa và báo cáo theo từng lớp.", "Snapshot thời tiết chỉ bao phủ một số điểm tham chiếu, nên cần mở rộng phạm vi và quy trình cập nhật provenance."])
    h2("9.3 Hướng phát triển")
    body("Ưu tiên tiếp theo là hoàn thiện dashboard Lab, bổ sung test tự động cho giao diện, thiết lập đánh giá mô hình và data card, đưa suy luận AI sang hàng đợi worker, bổ sung distributed scheduling, thử tải và kiểm thử bảo mật. Với cộng đồng, cần công bố release có checksum, duy trì SBOM theo mỗi phiên bản và dùng issue tracker để quản lý backlog theo module.")
    callout("KẾT LUẬN", "DX OS thể hiện đúng tư duy nền tảng: Core cung cấp dịch vụ dùng chung và giữ nguyên khi đổi ứng dụng; H P D I liên kết con người, quy trình, dữ liệu và trí tuệ; hai ứng dụng chứng minh giá trị sử dụng và khả năng mở rộng. Các giới hạn được công bố để Hội đồng phân biệt rõ phần đã kiểm chứng với hướng phát triển.", PALE_GREEN, TEAL)
    h2("Tài liệu tham chiếu nội bộ")
    bullets(["Đề thi PMNM HaUI 2026 bản chính thức, mục 3 đến 11.", "README dự án; ARCHITECTURE_OPEN_CORE; REPORT_OUTLINE; DEMO_SCRIPT; COUNCIL_QA.", "LICENSES; CONTRIBUTING; README của backend, app-farm, app-lab và ai.", "RUNBOOK_KIEM_THU và CV_RICE_IMAGE_VALIDATION_SUMMARY, cập nhật đến ngày 13 10 2026."])

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc.build(story)
    return PDF_OUT


if __name__ == "__main__":
    print(build_docx())
    print(build_pdf())
